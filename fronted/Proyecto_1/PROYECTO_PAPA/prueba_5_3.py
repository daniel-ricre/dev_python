import PySimpleGUI as sg
import pandas as pd
import os
from datetime import datetime
import calendar
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- Configuración de la ventana - TEMA MODERNO ---
sg.theme('DarkBlue3')  # Tema oscuro moderno
sg.set_options(font=('Segoe UI', 10))

# --- Colores personalizados ---
COLOR_PRIMARY = '#4A90E2'      # Azul moderno
COLOR_SUCCESS = '#27AE60'      # Verde éxito
COLOR_WARNING = '#F39C12'      # Naranja advertencia
COLOR_DANGER = '#E74C3C'       # Rojo peligro
COLOR_PURPLE = '#9B59B6'       # Púrpura
COLOR_DARK = '#2C3E50'         # Azul oscuro
COLOR_LIGHT = '#ECF0F1'        # Gris claro

# --- Archivo Excel con pestañas separadas ---
ARCHIVO_EXCEL = 'registro_inventario.xlsx'
HOJAS_EXCEL = {
    'Vino': 'VINO',
    'Latas': 'LATAS',
    'Dulces': 'DULCES'
}

# --- Configuración de etiquetas y colores por categoría ---
ETIQUETAS_CATEGORIA = {
    'Vino': {
        'nombre': '🍷 Nombre del vino:',
        'cantidad': '📦 Cantidad (botellas):',
        'titulo': '🍷 VINO',
        'color_btn': ('white', '#8B0000'),
        'color_titulo': '8B0000',
        'color_encabezado': 'CD5C5C',
        'icono': '🍷'
    },
    'Latas': {
        'nombre': '🥫 Nombre de la lata:',
        'cantidad': '📦 Cantidad (unidades):',
        'titulo': '🥫 LATAS',
        'color_btn': ('white', '#00008B'),
        'color_titulo': '00008B',
        'color_encabezado': '4682B4',
        'icono': '🥫'
    },
    'Dulces': {
        'nombre': '🍬 Nombre del dulce:',
        'cantidad': '📦 Cantidad (piezas):',
        'titulo': '🍬 DULCES',
        'color_btn': ('white', '#006400'),
        'color_titulo': '006400',
        'color_encabezado': '3CB371',
        'icono': '🍬'
    }
}

ALTURA_FILA_DATOS = 25

# --- Funciones de Excel (mantenidas igual) ---


def inicializar_excel():
    if not os.path.exists(ARCHIVO_EXCEL):
        from openpyxl import Workbook
        wb = Workbook()
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        for categoria, nombre_hoja in HOJAS_EXCEL.items():
            ws = wb.create_sheet(nombre_hoja)
            fila_actual = 1
            ws.merge_cells(f'A{fila_actual}:D{fila_actual}')
            celda_titulo = ws[f'A{fila_actual}']
            celda_titulo.value = ETIQUETAS_CATEGORIA[categoria]['titulo']
            celda_titulo.font = Font(size=14, bold=True, color='FFFFFF')
            celda_titulo.fill = PatternFill(
                start_color=ETIQUETAS_CATEGORIA[categoria]['color_titulo'],
                end_color=ETIQUETAS_CATEGORIA[categoria]['color_titulo'],
                fill_type='solid'
            )
            celda_titulo.alignment = Alignment(
                horizontal='center', vertical='center')
            ws.row_dimensions[fila_actual].height = 30
            fila_actual += 1
            encabezados = ['Nombre', 'Cantidad', 'Fecha', 'Hora']
            for col, encabezado in enumerate(encabezados, start=1):
                celda = ws.cell(row=fila_actual, column=col)
                celda.value = encabezado
                celda.font = Font(bold=True, color='FFFFFF', size=11)
                celda.fill = PatternFill(
                    start_color=ETIQUETAS_CATEGORIA[categoria]['color_encabezado'],
                    end_color=ETIQUETAS_CATEGORIA[categoria]['color_encabezado'],
                    fill_type='solid'
                )
                celda.alignment = Alignment(
                    horizontal='center', vertical='center')
                celda.border = Border(
                    left=Side(style='thin', color='000000'),
                    right=Side(style='thin', color='000000'),
                    top=Side(style='thin', color='000000'),
                    bottom=Side(style='thin', color='000000')
                )
            ws.row_dimensions[fila_actual].height = ALTURA_FILA_DATOS
            ws.column_dimensions['A'].width = 35
            ws.column_dimensions['B'].width = 12
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 12
        wb.save(ARCHIVO_EXCEL)
        return True
    return False


def guardar_en_excel(categoria, nombre, cantidad, fecha=None, hora=None):
    if fecha is None:
        fecha = datetime.now().strftime("%Y-%m-%d")
    if hora is None:
        hora = datetime.now().strftime("%H:%M:%S")
    try:
        inicializar_excel()
        wb = load_workbook(ARCHIVO_EXCEL)
        ws = wb[HOJAS_EXCEL[categoria]]
        fila_datos = 3
        while ws.cell(row=fila_datos, column=1).value is not None:
            fila_datos += 1
            if fila_datos > 1000:
                return False, "❌ Error: No se encontró espacio para guardar"
        ws.row_dimensions[fila_datos].height = ALTURA_FILA_DATOS
        ws.cell(row=fila_datos, column=1, value=nombre)
        ws.cell(row=fila_datos, column=2, value=cantidad)
        ws.cell(row=fila_datos, column=3, value=fecha)
        ws.cell(row=fila_datos, column=4, value=hora)
        for col in range(1, 5):
            celda = ws.cell(row=fila_datos, column=col)
            celda.alignment = Alignment(
                horizontal='left', vertical='center', wrap_text=True)
            celda.border = Border(
                left=Side(style='thin', color='000000'),
                right=Side(style='thin', color='000000'),
                top=Side(style='thin', color='000000'),
                bottom=Side(style='thin', color='000000')
            )
        if (fila_datos - 3) % 2 == 0:
            for col in range(1, 5):
                ws.cell(row=fila_datos, column=col).fill = PatternFill(
                    start_color='F5F5F5', end_color='F5F5F5', fill_type='solid'
                )
        wb.save(ARCHIVO_EXCEL)
        return True, f"✅ Registro guardado en {categoria}"
    except Exception as e:
        return False, f"❌ Error al guardar: {str(e)}"


def filtrar_por_fecha(categoria, fecha_objetivo):
    if not os.path.exists(ARCHIVO_EXCEL):
        return []
    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        hoja_nombre = HOJAS_EXCEL[categoria]
        if hoja_nombre not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[hoja_nombre]
        registros = []
        fecha_buscada = fecha_objetivo.strftime("%Y-%m-%d")
        fila_datos = 3
        while True:
            nombre = ws.cell(row=fila_datos, column=1).value
            if nombre is None:
                break
            cantidad = ws.cell(row=fila_datos, column=2).value
            fecha = ws.cell(row=fila_datos, column=3).value
            hora = ws.cell(row=fila_datos, column=4).value
            if fecha and fecha == fecha_buscada:
                registros.append(
                    [nombre, cantidad, fecha, hora if hora else ""])
            fila_datos += 1
        wb.close()
        registros.sort(key=lambda x: x[3] if x[3] else "", reverse=True)
        return registros
    except Exception as e:
        print(f"Error filtrando: {e}")
        return []


def mostrar_calendario(fecha_actual, fecha_seleccionada_anterior=None):
    año_actual = fecha_actual.year
    mes_actual = fecha_actual.month
    if fecha_seleccionada_anterior is None:
        fecha_seleccionada_anterior = fecha_actual
    cal = calendar.monthcalendar(año_actual, mes_actual)

    # Layout del calendario con mejor estilo
    layout_calendario = [
        [sg.Text(f"{calendar.month_name[mes_actual]} {año_actual}",
                 font=('Segoe UI', 18, 'bold'),
                 justification='center',
                 expand_x=True,
                 text_color='#4A90E2')],
        [sg.Text('')],
        [sg.Text(dia, size=(4, 1), justification='center',
                 font=('Segoe UI', 10, 'bold'), text_color='#7F8C8D')
         for dia in ['Lun', 'Mar', 'Mié', 'Jue', 'Vie', 'Sáb', 'Dom']],
    ]

    botones_calendario = {}
    for semana in cal:
        fila_botones = []
        for dia in semana:
            if dia == 0:
                fila_botones.append(sg.Text('', size=(4, 2)))
            else:
                fecha_boton = datetime(año_actual, mes_actual, dia)
                if fecha_boton.date() == fecha_seleccionada_anterior.date():
                    color_boton = ('black', '#F1C40F')  # Amarillo dorado
                elif (dia == datetime.now().day and mes_actual == datetime.now().month
                      and año_actual == datetime.now().year):
                    color_boton = ('white', '#27AE60')  # Verde
                else:
                    color_boton = ('#2C3E50', '#ECF0F1')  # Gris claro

                key_boton = f'-DIA_CAL_{dia}-'
                boton = sg.Button(str(dia),
                                  size=(4, 2),
                                  key=key_boton,
                                  button_color=color_boton,
                                  font=('Segoe UI', 10, 'bold'))
                fila_botones.append(boton)
                botones_calendario[key_boton] = boton
        layout_calendario.append(fila_botones)

    layout_calendario.extend([
        [sg.Text('')],
        [sg.Button('◀  Mes Anterior', size=(14, 1), key='-MES_ANTERIOR-',
                   button_color=('#2C3E50', '#BDC3C7'), font=('Segoe UI', 10)),
         sg.Button('📍 Hoy', size=(8, 1), key='-HOY-',
                   button_color=('white', '#4A90E2'), font=('Segoe UI', 10, 'bold')),
         sg.Button('Mes Siguiente  ▶', size=(14, 1), key='-MES_SIGUIENTE-',
                   button_color=('#2C3E50', '#BDC3C7'), font=('Segoe UI', 10))],
        [sg.Text('')],
        [sg.Text('', key='-INFO_SELECCION-', size=(30, 1), justification='center',
                 text_color='#4A90E2', font=('Segoe UI', 10, 'bold'))],
        [sg.Text('')],
        [sg.Button('❌ Cancelar', size=(12, 1), button_color=('#2C3E50', '#BDC3C7')),
         sg.Button('✅ Seleccionar', size=(12, 1), button_color=('white', '#4A90E2'),
                   key='-SELECCIONAR-', font=('Segoe UI', 10, 'bold'))]
    ])

    window_cal = sg.Window('📅 Seleccionar Fecha',
                           layout_calendario,
                           modal=True,
                           element_justification='center',
                           finalize=True)

    fecha_seleccionada = fecha_seleccionada_anterior
    mes_temporal = mes_actual
    año_temporal = año_actual

    if fecha_seleccionada:
        window_cal['-INFO_SELECCION-'].update(
            f"📌 Fecha seleccionada: {fecha_seleccionada.strftime('%d/%m/%Y')}"
        )

    while True:
        event_cal, values_cal = window_cal.read()
        if event_cal in (sg.WINDOW_CLOSED, '❌ Cancelar'):
            fecha_seleccionada = None
            break
        if event_cal == '-SELECCIONAR-':
            if fecha_seleccionada:
                break
            else:
                sg.popup('⚠️ Por favor selecciona un día primero', title='Aviso')
        if event_cal == '-HOY-':
            fecha_seleccionada = datetime.now()
            break
        if event_cal == '-MES_ANTERIOR-':
            mes_temporal -= 1
            if mes_temporal < 1:
                mes_temporal = 12
                año_temporal -= 1
            window_cal.close()
            fecha_temporal = datetime(año_temporal, mes_temporal, 1)
            fecha_seleccionada = mostrar_calendario(
                fecha_temporal, fecha_seleccionada)
            break
        if event_cal == '-MES_SIGUIENTE-':
            mes_temporal += 1
            if mes_temporal > 12:
                mes_temporal = 1
                año_temporal += 1
            window_cal.close()
            fecha_temporal = datetime(año_temporal, mes_temporal, 1)
            fecha_seleccionada = mostrar_calendario(
                fecha_temporal, fecha_seleccionada)
            break
        if event_cal.startswith('-DIA_CAL_'):
            dia = int(event_cal.replace('-DIA_CAL_', '').replace('-', ''))
            nueva_fecha = datetime(año_temporal, mes_temporal, dia)
            fecha_seleccionada = nueva_fecha
            for key, boton in botones_calendario.items():
                dia_boton = int(key.replace('-DIA_CAL_', '').replace('-', ''))
                fecha_boton = datetime(año_temporal, mes_temporal, dia_boton)
                if fecha_boton.date() == fecha_seleccionada.date():
                    boton.update(button_color=('black', '#F1C40F'))
                elif (dia_boton == datetime.now().day and mes_temporal == datetime.now().month
                      and año_temporal == datetime.now().year):
                    boton.update(button_color=('white', '#27AE60'))
                else:
                    boton.update(button_color=('#2C3E50', '#ECF0F1'))
            window_cal['-INFO_SELECCION-'].update(
                f"📌 Fecha seleccionada: {fecha_seleccionada.strftime('%d/%m/%Y')}"
            )

    window_cal.close()
    return fecha_seleccionada


# --- Inicializar ---
inicializar_excel()
fecha_actual = datetime.now()
hora_actual = fecha_actual.strftime("%H:%M:%S")
categoria_actual = 'Vino'

# --- INTERFAZ PRINCIPAL REDISEÑADA ---

# Cabecera con título
header = [
    [sg.Text('🏪 SISTEMA DE INVENTARIO', font=('Segoe UI', 20, 'bold'),
             text_color='#4A90E2', justification='center', expand_x=True)],
    [sg.Text('Gestión de productos por categoría', font=('Segoe UI', 11),
             text_color='#7F8C8D', justification='center', expand_x=True)],
    [sg.Text('')]
]

# Selector de categoría con botones más atractivos
categoria_selector = [
    [sg.Text('📂 SELECCIONAR CATEGORÍA', font=('Segoe UI', 12, 'bold'),
             text_color='#ECF0F1', justification='center', expand_x=True)],
    [sg.Button('🍷  VINO', size=(15, 2), font=('Segoe UI', 12, 'bold'),
               button_color=ETIQUETAS_CATEGORIA['Vino']['color_btn'],
               key='-CAT_VINO-', border_width=2),
     sg.Button('🥫  LATAS', size=(15, 2), font=('Segoe UI', 12, 'bold'),
               button_color=ETIQUETAS_CATEGORIA['Latas']['color_btn'],
               key='-CAT_LATAS-', border_width=2),
     sg.Button('🍬  DULCES', size=(15, 2), font=('Segoe UI', 12, 'bold'),
               button_color=ETIQUETAS_CATEGORIA['Dulces']['color_btn'],
               key='-CAT_DULCES-', border_width=2)],
    [sg.Text(f'📍 Pestaña activa: {HOJAS_EXCEL[categoria_actual]}',
             key='-PESTANA_ACTIVA-', font=('Segoe UI', 10, 'italic'),
             text_color='#4A90E2', justification='center', expand_x=True)]
]

# Selector de fecha
fecha_selector = [
    [sg.Text('')],
    [sg.Button('📅  SELECCIONAR FECHA', size=(22, 2), font=('Segoe UI', 11, 'bold'),
               button_color=('white', '#9B59B6'), key='-SELECCIONAR_FECHA-', border_width=2),
     sg.Text(fecha_actual.strftime("%d de %B de %Y"), key='-FECHA_MOSTRADA-',
             font=('Segoe UI', 13, 'bold'), text_color='#F1C40F',
             size=(30, 1), justification='center')]
]

# Tabla de registros
tabla_frame = [
    [sg.Text('📋 REGISTROS DEL DÍA', font=('Segoe UI', 12, 'bold'),
             text_color='#ECF0F1', key='-TITULO_TABLA-')],
    [sg.Table(
        values=[],
        headings=['Nombre', 'Cantidad', 'Fecha', 'Hora'],
        max_col_width=30,
        auto_size_columns=False,
        col_widths=[28, 10, 13, 10],
        display_row_numbers=True,
        justification='left',
        num_rows=18,
        key='-TABLA-',
        row_height=28,
        enable_events=True,
        expand_x=True,
        expand_y=True,
        font=('Segoe UI', 10),
        header_font=('Segoe UI', 10, 'bold'),
        header_text_color='#2C3E50',
        header_background_color='#BDC3C7',
        alternating_row_color='#34495E'
    )],
    [sg.Button('🔄  Actualizar', size=(14, 1), font=('Segoe UI', 10),
               button_color=('#2C3E50', '#BDC3C7')),
     sg.Button('📂  Abrir Excel', size=(14, 1), font=('Segoe UI', 10),
               button_color=('white', '#27AE60')),
     sg.Button('📊  Ver Todos', size=(14, 1), font=('Segoe UI', 10),
               button_color=('white', '#F39C12'))]
]

# Formulario de registro
formulario_frame = [
    [sg.Text(f'✨ NUEVO REGISTRO - {categoria_actual.upper()}',
             font=('Segoe UI', 13, 'bold'), text_color='#4A90E2',
             key='-TITULO_FORMULARIO-')],
    [sg.Text('')],
    [sg.Text(ETIQUETAS_CATEGORIA[categoria_actual]['nombre'],
             size=(22, 1), key='-LABEL_NOMBRE-', font=('Segoe UI', 11),
             text_color='#ECF0F1')],
    [sg.InputText(key='-NOMBRE-', size=(32, 1), font=('Segoe UI', 11),
                  border_width=2)],
    [sg.Text('')],
    [sg.Text(ETIQUETAS_CATEGORIA[categoria_actual]['cantidad'],
             size=(22, 1), key='-LABEL_CANTIDAD-', font=('Segoe UI', 11),
             text_color='#ECF0F1')],
    [sg.InputText(key='-CANTIDAD-', size=(32, 1), font=('Segoe UI', 11),
                  border_width=2)],
    [sg.Text('')],
    [sg.Text('📅 Fecha de registro:', size=(22, 1), font=('Segoe UI', 11),
             text_color='#ECF0F1')],
    [sg.InputText(key='-FECHA_REGISTRO-', size=(32, 1),
                  default_text=fecha_actual.strftime("%Y-%m-%d"),
                  font=('Segoe UI', 10), border_width=2)],
    [sg.Text('')],
    [sg.Text('🕐 Hora de registro:', size=(22, 1), font=('Segoe UI', 11),
             text_color='#ECF0F1')],
    [sg.InputText(key='-HORA_REGISTRO-', size=(32, 1),
                  default_text=hora_actual,
                  font=('Segoe UI', 10), border_width=2)],
    [sg.Text('')],
    [sg.Button('💾  GUARDAR REGISTRO', size=(22, 2),
               button_color=('white', '#4A90E2'),
               font=('Segoe UI', 12, 'bold'),
               bind_return_key=True,
               key='-GUARDAR-', border_width=3)],
    [sg.Text('')],
    [sg.Text('', key='-MENSAJE-', size=(40, 2), text_color='#27AE60',
             font=('Segoe UI', 10, 'bold'), justification='center')],
    [sg.Text('')],
    [sg.Button('🗑️  Limpiar campos', size=(16, 1), font=('Segoe UI', 10),
               button_color=('#2C3E50', '#BDC3C7')),
     sg.Button('🕐  Hora actual', size=(16, 1), font=('Segoe UI', 10),
               button_color=('white', '#3498DB'), key='-HORA_ACTUAL-')]
]

# Layout principal con columnas
layout = header + [
    [sg.HorizontalSeparator(color='#4A90E2')],
    *categoria_selector,
    [sg.HorizontalSeparator(color='#34495E')],
    *fecha_selector,
    [sg.HorizontalSeparator(color='#34495E')],
    [sg.Text('')],
    [
        sg.Column(tabla_frame, expand_x=True, expand_y=True,
                  element_justification='left'),
        sg.VSeparator(color='#4A90E2'),
        sg.Column(formulario_frame, element_justification='center',
                  vertical_alignment='top')
    ],
    [sg.HorizontalSeparator(color='#4A90E2')],
    [sg.Text('💡 Tip: Usa el calendario para filtrar por fecha. Cada categoría tiene su propia pestaña en Excel.',
             text_color='#7F8C8D', font=('Segoe UI', 9, 'italic'), expand_x=True)]
]

# Crear ventana
window = sg.Window('🏪 Tienda Registro - Sistema de Inventario Profesional',
                   layout,
                   size=(1150, 720),
                   resizable=True,
                   finalize=True,
                   icon=None)

# Variables globales
fecha_seleccionada_global = fecha_actual


def actualizar_tabla_por_fecha(categoria, fecha_filtro):
    registros = filtrar_por_fecha(categoria, fecha_filtro)
    window['-TABLA-'].update(values=registros)
    fecha_formateada = fecha_filtro.strftime("%d/%m/%Y")
    window['-TITULO_TABLA-'].update(
        f'📋 {categoria.upper()} - Registros del {fecha_formateada}')
    return len(registros) > 0


def ver_todos_registros(categoria):
    if not os.path.exists(ARCHIVO_EXCEL):
        return False
    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        ws = wb[HOJAS_EXCEL[categoria]]
        registros = []
        fila_datos = 3
        while True:
            nombre = ws.cell(row=fila_datos, column=1).value
            if nombre is None:
                break
            cantidad = ws.cell(row=fila_datos, column=2).value
            fecha = ws.cell(row=fila_datos, column=3).value
            hora = ws.cell(row=fila_datos, column=4).value
            registros.append([nombre, cantidad, fecha, hora if hora else ""])
            fila_datos += 1
        wb.close()
        if registros:
            registros.sort(key=lambda x: f"{x[2]} {x[3]}", reverse=True)
            window['-TABLA-'].update(values=registros)
            window['-TITULO_TABLA-'].update(
                f'📋 {categoria.upper()} - TODOS LOS REGISTROS')
            return True
    except Exception as e:
        sg.popup_error(f"Error al cargar: {e}")
    return False


def cambiar_categoria(nueva_categoria):
    global categoria_actual
    categoria_actual = nueva_categoria
    window['-LABEL_NOMBRE-'].update(
        ETIQUETAS_CATEGORIA[categoria_actual]['nombre'])
    window['-LABEL_CANTIDAD-'].update(
        ETIQUETAS_CATEGORIA[categoria_actual]['cantidad'])
    window['-TITULO_FORMULARIO-'].update(
        f'✨ NUEVO REGISTRO - {categoria_actual.upper()}')
    window['-PESTANA_ACTIVA-'].update(
        f'📍 Pestaña activa: {HOJAS_EXCEL[categoria_actual]}')
    actualizar_tabla_por_fecha(categoria_actual, fecha_seleccionada_global)
    window['-MENSAJE-'].update(
        f'✅ Categoría cambiada a {categoria_actual}', text_color='#27AE60')


# Cargar datos iniciales
actualizar_tabla_por_fecha(categoria_actual, fecha_actual)

# Bucle principal
while True:
    event, values = window.read()

    if event == sg.WINDOW_CLOSED:
        break

    window['-FECHA_MOSTRADA-'].update(
        fecha_seleccionada_global.strftime("%d de %B de %Y"))

    if event == '-CAT_VINO-':
        cambiar_categoria('Vino')
    elif event == '-CAT_LATAS-':
        cambiar_categoria('Latas')
    elif event == '-CAT_DULCES-':
        cambiar_categoria('Dulces')

    if event == '-SELECCIONAR_FECHA-':
        fecha_elegida = mostrar_calendario(
            fecha_actual, fecha_seleccionada_global)
        if fecha_elegida:
            fecha_seleccionada_global = fecha_elegida
            window['-FECHA_REGISTRO-'].update(
                fecha_elegida.strftime("%Y-%m-%d"))
            actualizar_tabla_por_fecha(
                categoria_actual, fecha_seleccionada_global)
            window['-MENSAJE-'].update(f"✅ Mostrando registros del {fecha_elegida.strftime('%d/%m/%Y')}",
                                       text_color='#27AE60')

    if event == '-HORA_ACTUAL-':
        hora_actual = datetime.now().strftime("%H:%M:%S")
        window['-HORA_REGISTRO-'].update(hora_actual)
        window['-MENSAJE-'].update("🕐 Hora actual actualizada",
                                   text_color='#3498DB')

    if event == '-GUARDAR-':
        nombre = values['-NOMBRE-'].strip()
        cantidad = values['-CANTIDAD-'].strip()
        fecha = values['-FECHA_REGISTRO-'].strip()
        hora = values['-HORA_REGISTRO-'].strip()

        if not nombre or not cantidad or not fecha or not hora:
            window['-MENSAJE-'].update(
                "❌ Todos los campos son obligatorios", text_color='#E74C3C')
        elif not cantidad.isdigit():
            window['-MENSAJE-'].update("❌ La cantidad debe ser un número",
                                       text_color='#E74C3C')
        else:
            exito, mensaje = guardar_en_excel(
                categoria_actual, nombre, int(cantidad), fecha, hora)
            if exito:
                window['-MENSAJE-'].update(mensaje, text_color='#27AE60')
                window['-NOMBRE-'].update('')
                window['-CANTIDAD-'].update('')
                actualizar_tabla_por_fecha(
                    categoria_actual, fecha_seleccionada_global)
                window['-NOMBRE-'].set_focus()
            else:
                window['-MENSAJE-'].update(mensaje, text_color='#E74C3C')

    if event == '🔄  Actualizar':
        actualizar_tabla_por_fecha(categoria_actual, fecha_seleccionada_global)
        window['-MENSAJE-'].update("✅ Tabla actualizada", text_color='#27AE60')

    if event == '📊  Ver Todos':
        if ver_todos_registros(categoria_actual):
            window['-MENSAJE-'].update(f"✅ Mostrando todos los registros de {categoria_actual}",
                                       text_color='#27AE60')
        else:
            window['-MENSAJE-'].update(f"ℹ️ No hay registros en {categoria_actual}",
                                       text_color='#F39C12')

    if event == '📂  Abrir Excel':
        if os.path.exists(ARCHIVO_EXCEL):
            os.startfile(ARCHIVO_EXCEL)
        else:
            sg.popup("📁 Aún no hay archivo Excel. Guarda tu primer registro.",
                     title="Información")

    if event == '🗑️  Limpiar campos':
        window['-NOMBRE-'].update('')
        window['-CANTIDAD-'].update('')
        window['-MENSAJE-'].update("🧹 Campos limpiados", text_color='#3498DB')
        window['-NOMBRE-'].set_focus()

window.close()
