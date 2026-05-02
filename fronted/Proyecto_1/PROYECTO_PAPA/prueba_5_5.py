import PySimpleGUI as sg
import pandas as pd
import os
from datetime import datetime
import calendar
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- Tema minimalista moderno ---
sg.theme('LightGray1')
sg.set_options(font=('Segoe UI', 10))

# --- Paleta de colores corporativa ---
COLORS = {
    'primary': '#6366F1',
    'success': '#10B981',
    'warning': '#F59E0B',
    'danger': '#EF4444',
    'dark': '#1E293B',
    'light': '#F8FAFC',
    'card': '#FFFFFF',
    'border': '#E2E8F0',
    'text': '#334155',
    'text_light': '#64748B',
    'accent': '#8B5CF6'
}

# --- Archivo Excel ---
ARCHIVO_EXCEL = 'registro_inventario.xlsx'
HOJAS_EXCEL = {
    'Vino': 'VINO',
    'Latas': 'LATAS',
    'Dulces': 'DULCES'
}

# --- Configuración por categoría ---
CATEGORIAS = {
    'Vino': {
        'nombre': 'Nombre del vino',
        'cantidad': 'Cantidad (botellas)',
        'icono': '🍷',
        'color': '#DC2626',
    },
    'Latas': {
        'nombre': 'Nombre de la lata',
        'cantidad': 'Cantidad (unidades)',
        'icono': '🥫',
        'color': '#2563EB',
    },
    'Dulces': {
        'nombre': 'Nombre del dulce',
        'cantidad': 'Cantidad (piezas)',
        'icono': '🍬',
        'color': '#059669',
    }
}

ALTURA_FILA_DATOS = 25

# --- Funciones de Excel ---


def inicializar_excel():
    if not os.path.exists(ARCHIVO_EXCEL):
        from openpyxl import Workbook
        wb = Workbook()
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        for categoria, nombre_hoja in HOJAS_EXCEL.items():
            ws = wb.create_sheet(nombre_hoja)
            fila_actual = 1
            ws.merge_cells(f'A{fila_actual}:D{fila_actual}')
            celda_titulo = ws[f'A{fila_actual}']
            celda_titulo.value = f"{CATEGORIAS[categoria]['icono']} {categoria.upper()}"
            celda_titulo.font = Font(size=14, bold=True, color='FFFFFF')
            color_hex = CATEGORIAS[categoria]['color'].replace('#', '')
            celda_titulo.fill = PatternFill(
                start_color=color_hex, end_color=color_hex, fill_type='solid')
            celda_titulo.alignment = Alignment(
                horizontal='center', vertical='center')
            ws.row_dimensions[fila_actual].height = 30
            fila_actual += 1
            encabezados = ['Nombre', 'Cantidad', 'Fecha', 'Hora']
            for col, encabezado in enumerate(encabezados, start=1):
                celda = ws.cell(row=fila_actual, column=col)
                celda.value = encabezado
                celda.font = Font(bold=True, color='FFFFFF', size=11)
                celda.fill = PatternFill(
                    start_color='475569', end_color='475569', fill_type='solid')
                celda.alignment = Alignment(
                    horizontal='center', vertical='center')
            ws.row_dimensions[fila_actual].height = ALTURA_FILA_DATOS
            ws.column_dimensions['A'].width = 35
            ws.column_dimensions['B'].width = 12
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 12
        wb.save(ARCHIVO_EXCEL)
        return True
    return False


def guardar_en_excel(categoria, nombre, cantidad, fecha=None, hora=None):
    if fecha is None:
        fecha = datetime.now().strftime("%Y-%m-%d")
    if hora is None:
        hora = datetime.now().strftime("%H:%M:%S")
    try:
        inicializar_excel()
        wb = load_workbook(ARCHIVO_EXCEL)
        ws = wb[HOJAS_EXCEL[categoria]]
        fila_datos = 3
        while ws.cell(row=fila_datos, column=1).value is not None:
            fila_datos += 1
        ws.row_dimensions[fila_datos].height = ALTURA_FILA_DATOS
        ws.cell(row=fila_datos, column=1, value=nombre)
        ws.cell(row=fila_datos, column=2, value=cantidad)
        ws.cell(row=fila_datos, column=3, value=fecha)
        ws.cell(row=fila_datos, column=4, value=hora)
        wb.save(ARCHIVO_EXCEL)
        return True, f"✅ Registro guardado en {categoria}"
    except Exception as e:
        return False, f"❌ Error: {str(e)}"


def filtrar_por_fecha(categoria, fecha_objetivo):
    if not os.path.exists(ARCHIVO_EXCEL):
        return []
    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        ws = wb[HOJAS_EXCEL[categoria]]
        registros = []
        fecha_buscada = fecha_objetivo.strftime("%Y-%m-%d")
        fila_datos = 3
        while True:
            nombre = ws.cell(row=fila_datos, column=1).value
            if nombre is None:
                break
            cantidad = ws.cell(row=fila_datos, column=2).value
            fecha = ws.cell(row=fila_datos, column=3).value
            hora = ws.cell(row=fila_datos, column=4).value
            if fecha and fecha == fecha_buscada:
                registros.append(
                    [nombre, cantidad, fecha, hora if hora else ""])
            fila_datos += 1
        wb.close()
        registros.sort(key=lambda x: x[3] if x[3] else "", reverse=True)
        return registros
    except:
        return []


def mostrar_calendario(fecha_actual, fecha_seleccionada_anterior=None):
    if fecha_seleccionada_anterior is None:
        fecha_seleccionada_anterior = fecha_actual

    año, mes = fecha_actual.year, fecha_actual.month
    cal = calendar.monthcalendar(año, mes)

    layout_cal = [
        [sg.Text(f"{calendar.month_name[mes]} {año}",
                 font=('Segoe UI', 16, 'bold'),
                 text_color=COLORS['primary'],
                 justification='center', expand_x=True)],
        [sg.Text('')],
        [sg.Text(d, size=(4, 1), justification='center',
                 font=('Segoe UI', 9, 'bold'), text_color=COLORS['text_light'])
         for d in ['L', 'M', 'X', 'J', 'V', 'S', 'D']],
    ]

    botones = {}
    for semana in cal:
        fila = []
        for dia in semana:
            if dia == 0:
                fila.append(sg.Text('', size=(4, 2)))
            else:
                fecha_btn = datetime(año, mes, dia)
                if fecha_btn.date() == fecha_seleccionada_anterior.date():
                    color = ('white', COLORS['primary'])
                elif fecha_btn.date() == datetime.now().date():
                    color = (COLORS['dark'], COLORS['warning'])
                else:
                    color = (COLORS['text'], COLORS['light'])

                key = f'-DIA_{dia}-'
                btn = sg.Button(str(dia), size=(4, 2), key=key, button_color=color,
                                font=('Segoe UI', 9, 'bold'), border_width=0)
                fila.append(btn)
                botones[key] = btn
        layout_cal.append(fila)

    layout_cal.extend([
        [sg.Text('')],
        [sg.Button('◀', size=(4, 1), key='-MES_ANT-', button_color=(COLORS['text'], COLORS['light']), border_width=0),
         sg.Button('Hoy', size=(6, 1), key='-HOY-',
                   button_color=('white', COLORS['primary']), border_width=0),
         sg.Button('▶', size=(4, 1), key='-MES_SIG-', button_color=(COLORS['text'], COLORS['light']), border_width=0)],
        [sg.Text('')],
        [sg.Button('Cancelar', size=(10, 1), button_color=(COLORS['text'], COLORS['light']), border_width=0),
         sg.Button('Seleccionar', size=(10, 1), button_color=('white', COLORS['success']), border_width=0, key='-SEL-')]
    ])

    win = sg.Window('📅 Calendario', layout_cal, modal=True,
                    element_justification='center', finalize=True)

    fecha_sel = fecha_seleccionada_anterior
    mes_temp, año_temp = mes, año

    while True:
        ev, _ = win.read()
        if ev in (sg.WINDOW_CLOSED, 'Cancelar'):
            fecha_sel = None
            break
        if ev == '-SEL-':
            break
        if ev == '-HOY-':
            fecha_sel = datetime.now()
            break
        if ev == '-MES_ANT-':
            mes_temp -= 1
            if mes_temp < 1:
                mes_temp = 12
                año_temp -= 1
            win.close()
            fecha_sel = mostrar_calendario(
                datetime(año_temp, mes_temp, 1), fecha_sel)
            break
        if ev == '-MES_SIG-':
            mes_temp += 1
            if mes_temp > 12:
                mes_temp = 1
                año_temp += 1
            win.close()
            fecha_sel = mostrar_calendario(
                datetime(año_temp, mes_temp, 1), fecha_sel)
            break
        if ev.startswith('-DIA_'):
            dia = int(ev.replace('-DIA_', '').replace('-', ''))
            fecha_sel = datetime(año_temp, mes_temp, dia)
            for k, btn in botones.items():
                d = int(k.replace('-DIA_', '').replace('-', ''))
                f = datetime(año_temp, mes_temp, d)
                if f.date() == fecha_sel.date():
                    btn.update(button_color=('white', COLORS['primary']))
                elif f.date() == datetime.now().date():
                    btn.update(button_color=(
                        COLORS['dark'], COLORS['warning']))
                else:
                    btn.update(button_color=(COLORS['text'], COLORS['light']))

    win.close()
    return fecha_sel


# --- Inicialización ---
inicializar_excel()
fecha_actual = datetime.now()
hora_actual = fecha_actual.strftime("%H:%M")
categoria_actual = 'Vino'

# --- Función para crear botón de categoría ---


def cat_button(text, key, color):
    return sg.Button(text, size=(14, 2), key=key, button_color=('white', color),
                     font=('Segoe UI', 11, 'bold'), border_width=0, pad=(5, 5))

# --- INTERFAZ GRÁFICA CORREGIDA ---


# Cabecera
header = [
    [sg.Text('🏪', font=('Segoe UI', 28), justification='center', expand_x=True)],
    [sg.Text('INVENTARIO PRO', font=('Segoe UI', 18, 'bold'),
             text_color=COLORS['dark'], justification='center', expand_x=True)],
    [sg.Text('Gestión inteligente de productos', font=('Segoe UI', 10),
             text_color=COLORS['text_light'], justification='center', expand_x=True)],
]

# Selector de categoría
categoria_layout = [
    [sg.Text('📂 CATEGORÍAS', font=('Segoe UI', 11, 'bold'),
             text_color=COLORS['text_light'])],
    [cat_button(f"{CATEGORIAS['Vino']['icono']}  VINO", '-CAT_VINO-', CATEGORIAS['Vino']['color']),
     cat_button(f"{CATEGORIAS['Latas']['icono']}  LATAS",
                '-CAT_LATAS-', CATEGORIAS['Latas']['color']),
     cat_button(f"{CATEGORIAS['Dulces']['icono']}  DULCES", '-CAT_DULCES-', CATEGORIAS['Dulces']['color'])],
    [sg.Text(f"📍 Pestaña activa: {HOJAS_EXCEL[categoria_actual]}",
             key='-PESTANA_ACTIVA-', font=('Segoe UI', 9, 'italic'),
             text_color=COLORS['primary'], justification='center', expand_x=True)]
]

# Selector de fecha
fecha_layout = [
    [sg.Text('📅 FECHA DE CONSULTA', font=('Segoe UI', 11, 'bold'),
             text_color=COLORS['text_light'])],
    [sg.Button('📆  SELECCIONAR FECHA', size=(20, 2), font=('Segoe UI', 10, 'bold'),
               button_color=('white', COLORS['accent']), key='-SELECCIONAR_FECHA-', border_width=0),
     sg.Text(fecha_actual.strftime("%d de %B de %Y"), key='-FECHA_MOSTRADA-',
             font=('Segoe UI', 13, 'bold'), text_color=COLORS['dark'],
             size=(25, 1), justification='center')]
]

# Tabla de registros
tabla_layout = [
    [sg.Text('📋 REGISTROS DEL DÍA', font=('Segoe UI', 12, 'bold'),
             text_color=COLORS['dark'], key='-TITULO_TABLA-')],
    [sg.Table(
        values=[],
        headings=['Producto', 'Cant', 'Fecha', 'Hora'],
        col_widths=[28, 8, 13, 10],
        display_row_numbers=True,
        num_rows=16,
        key='-TABLA-',
        row_height=28,
        font=('Segoe UI', 9),
        header_font=('Segoe UI', 9, 'bold'),
        header_text_color=COLORS['light'],
        header_background_color=COLORS['dark'],
        alternating_row_color='#F1F5F9',
        expand_x=True, expand_y=True
    )],
    [sg.Button('🔄 Actualizar', size=(12, 1), font=('Segoe UI', 9),
               button_color=(COLORS['text'], COLORS['light']), border_width=0),
     sg.Button('📂 Abrir Excel', size=(12, 1), font=('Segoe UI', 9),
               button_color=('white', COLORS['success']), border_width=0),
     sg.Button('📊 Ver todos', size=(12, 1), font=('Segoe UI', 9),
               button_color=('white', COLORS['warning']), border_width=0)]
]

# Formulario (con elementos que se actualizarán vía update)
cat = CATEGORIAS[categoria_actual]
formulario_layout = [
    [sg.Text(f"{cat['icono']}  NUEVO REGISTRO", font=('Segoe UI', 13, 'bold'),
             text_color=cat['color'], key='-TITULO_FORMULARIO-')],
    [sg.Text('')],
    [sg.Text(cat['nombre'], font=('Segoe UI', 10, 'bold'),
             text_color=COLORS['text'], key='-LABEL_NOMBRE-')],
    [sg.Input(key='-NOMBRE-', size=(28, 1), font=('Segoe UI', 11))],
    [sg.Text('')],
    [sg.Text(cat['cantidad'], font=('Segoe UI', 10, 'bold'),
             text_color=COLORS['text'], key='-LABEL_CANTIDAD-')],
    [sg.Input(key='-CANTIDAD-', size=(28, 1), font=('Segoe UI', 11))],
    [sg.Text('')],
    [sg.Text('📅 Fecha', font=('Segoe UI', 10, 'bold'), text_color=COLORS['text'])],
    [sg.Input(key='-FECHA_REGISTRO-', size=(28, 1),
              default_text=fecha_actual.strftime("%Y-%m-%d"), font=('Segoe UI', 10))],
    [sg.Text('')],
    [sg.Text('🕐 Hora', font=('Segoe UI', 10, 'bold'), text_color=COLORS['text'])],
    [sg.Input(key='-HORA_REGISTRO-', size=(28, 1),
              default_text=hora_actual, font=('Segoe UI', 10))],
    [sg.Text('')],
    [sg.Button('💾  GUARDAR', size=(20, 2), button_color=('white', COLORS['primary']),
               font=('Segoe UI', 11, 'bold'), key='-GUARDAR-', border_width=0)],
    [sg.Text('')],
    [sg.Text('', key='-MENSAJE-', size=(35, 2), text_color=COLORS['success'],
             font=('Segoe UI', 9, 'bold'), justification='center')],
    [sg.Button('🗑️ Limpiar', size=(12, 1), font=('Segoe UI', 9),
               button_color=(COLORS['text'], COLORS['light']), border_width=0),
     sg.Button('🕐 Ahora', size=(12, 1), font=('Segoe UI', 9),
               button_color=('white', COLORS['primary']), key='-HORA_ACTUAL-', border_width=0)]
]

# Layout principal
layout = [
    header,
    [sg.Text('')],
    [sg.Frame('', categoria_layout, relief=sg.RELIEF_FLAT, border_width=0,
              element_justification='center', expand_x=True)],
    [sg.Text('')],
    [sg.Frame('', fecha_layout, relief=sg.RELIEF_FLAT, border_width=0,
              element_justification='center', expand_x=True)],
    [sg.Text('')],
    [
        sg.Frame('', tabla_layout, relief=sg.RELIEF_FLAT, border_width=1,
                 border_color=COLORS['border'], expand_x=True, expand_y=True),
        sg.Column([[sg.VSeparator(color=COLORS['border'])]]),
        sg.Frame('', formulario_layout, relief=sg.RELIEF_FLAT, border_width=1,
                 border_color=COLORS['border'])
    ],
    [sg.Text('')],
    [sg.Text('💡 Selecciona una categoría y fecha para gestionar tu inventario',
             text_color=COLORS['text_light'], font=('Segoe UI', 8, 'italic'),
             justification='center', expand_x=True)]
]

# Crear ventana
window = sg.Window('🏪 Inventario Pro', layout, size=(1100, 700),
                   resizable=True, finalize=True, margins=(15, 15))

# Variables globales
fecha_seleccionada_global = fecha_actual


def actualizar_tabla(categoria, fecha):
    regs = filtrar_por_fecha(categoria, fecha)
    window['-TABLA-'].update(values=regs)
    window['-TITULO_TABLA-'].update(
        f'📋 {categoria.upper()} - {fecha.strftime("%d/%m/%Y")}')
    return len(regs) > 0


def ver_todos(categoria):
    if not os.path.exists(ARCHIVO_EXCEL):
        return False
    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        ws = wb[HOJAS_EXCEL[categoria]]
        regs = []
        fila = 3
        while True:
            n = ws.cell(row=fila, column=1).value
            if n is None:
                break
            regs.append([n, ws.cell(row=fila, column=2).value,
                        ws.cell(row=fila, column=3).value,
                        ws.cell(row=fila, column=4).value or ""])
            fila += 1
        wb.close()
        if regs:
            regs.sort(key=lambda x: f"{x[2]} {x[3]}", reverse=True)
            window['-TABLA-'].update(values=regs)
            window['-TITULO_TABLA-'].update(f'📋 {categoria.upper()} - TODOS')
            return True
    except:
        pass
    return False


def cambiar_categoria(nueva_cat):
    global categoria_actual
    categoria_actual = nueva_cat
    cat = CATEGORIAS[categoria_actual]

    # Actualizar textos en lugar de reconstruir
    window['-TITULO_FORMULARIO-'].update(f"{cat['icono']}  NUEVO REGISTRO")
    window['-TITULO_FORMULARIO-'].update(text_color=cat['color'])
    window['-LABEL_NOMBRE-'].update(cat['nombre'])
    window['-LABEL_CANTIDAD-'].update(cat['cantidad'])
    window['-PESTANA_ACTIVA-'].update(
        f'📍 Pestaña activa: {HOJAS_EXCEL[categoria_actual]}')

    actualizar_tabla(categoria_actual, fecha_seleccionada_global)
    window['-MENSAJE-'].update(
        f'✅ Categoría: {categoria_actual}', text_color=COLORS['success'])


# Cargar inicial
actualizar_tabla(categoria_actual, fecha_actual)

# Bucle principal
while True:
    event, values = window.read()

    if event == sg.WINDOW_CLOSED:
        break

    window['-FECHA_MOSTRADA-'].update(
        fecha_seleccionada_global.strftime("%d de %B de %Y"))

    if event == '-CAT_VINO-':
        cambiar_categoria('Vino')
    elif event == '-CAT_LATAS-':
        cambiar_categoria('Latas')
    elif event == '-CAT_DULCES-':
        cambiar_categoria('Dulces')

    if event == '-SELECCIONAR_FECHA-':
        nueva_fecha = mostrar_calendario(
            fecha_actual, fecha_seleccionada_global)
        if nueva_fecha:
            fecha_seleccionada_global = nueva_fecha
            window['-FECHA_REGISTRO-'].update(nueva_fecha.strftime("%Y-%m-%d"))
            actualizar_tabla(categoria_actual, nueva_fecha)
            window['-MENSAJE-'].update(f'✅ Fecha: {nueva_fecha.strftime("%d/%m/%Y")}',
                                       text_color=COLORS['success'])

    if event == '-HORA_ACTUAL-':
        window['-HORA_REGISTRO-'].update(datetime.now().strftime("%H:%M"))
        window['-MENSAJE-'].update('🕐 Hora actualizada',
                                   text_color=COLORS['primary'])

    if event == '-GUARDAR-':
        nombre = values['-NOMBRE-'].strip()
        cantidad = values['-CANTIDAD-'].strip()
        fecha = values['-FECHA_REGISTRO-'].strip()
        hora = values['-HORA_REGISTRO-'].strip()

        if not all([nombre, cantidad, fecha, hora]):
            window['-MENSAJE-'].update('❌ Todos los campos obligatorios',
                                       text_color=COLORS['danger'])
        elif not cantidad.isdigit():
            window['-MENSAJE-'].update('❌ Cantidad debe ser número',
                                       text_color=COLORS['danger'])
        else:
            exito, msg = guardar_en_excel(
                categoria_actual, nombre, int(cantidad), fecha, hora)
            if exito:
                window['-MENSAJE-'].update(msg, text_color=COLORS['success'])
                window['-NOMBRE-'].update('')
                window['-CANTIDAD-'].update('')
                actualizar_tabla(categoria_actual, fecha_seleccionada_global)
                window['-NOMBRE-'].set_focus()
            else:
                window['-MENSAJE-'].update(msg, text_color=COLORS['danger'])

    if event == '🔄 Actualizar':
        actualizar_tabla(categoria_actual, fecha_seleccionada_global)
        window['-MENSAJE-'].update('✅ Tabla actualizada',
                                   text_color=COLORS['success'])

    if event == '📊 Ver todos':
        if ver_todos(categoria_actual):
            window['-MENSAJE-'].update(f'✅ Mostrando todos los registros',
                                       text_color=COLORS['success'])
        else:
            window['-MENSAJE-'].update('ℹ️ No hay registros',
                                       text_color=COLORS['warning'])

    if event == '📂 Abrir Excel':
        if os.path.exists(ARCHIVO_EXCEL):
            os.startfile(ARCHIVO_EXCEL)

    if event == '🗑️ Limpiar':
        window['-NOMBRE-'].update('')
        window['-CANTIDAD-'].update('')
        window['-MENSAJE-'].update('🧹 Campos limpiados',
                                   text_color=COLORS['text_light'])
        window['-NOMBRE-'].set_focus()

window.close()
