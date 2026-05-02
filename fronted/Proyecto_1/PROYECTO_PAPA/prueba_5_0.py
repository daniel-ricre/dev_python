import PySimpleGUI as sg
import pandas as pd
import os
from datetime import datetime
import calendar
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- Configuración de la ventana ---
sg.theme('LightBlue3')

# --- ÚNICO archivo Excel ---
ARCHIVO_EXCEL = 'registro_inventario.xlsx'
NOMBRE_HOJA = 'INVENTARIO'

# --- Configuración de etiquetas dinámicas por categoría ---
ETIQUETAS_CATEGORIA = {
    'Vino': {
        'nombre': 'Nombre del vino:',
        'cantidad': 'Cantidad (botellas):',
        'titulo': '🍷 VINO',
        'color_titulo': '8B0000',  # Rojo oscuro
        'color_encabezado': 'CD5C5C'  # Rojo indio
    },
    'Latas': {
        'nombre': 'Nombre de la lata:',
        'cantidad': 'Cantidad (unidades):',
        'titulo': '🥫 LATAS',
        'color_titulo': '00008B',  # Azul oscuro
        'color_encabezado': '4682B4'  # Azul acero
    },
    'Dulces': {
        'nombre': 'Nombre del dulce:',
        'cantidad': 'Cantidad (piezas):',
        'titulo': '🍬 DULCES',
        'color_titulo': '006400',  # Verde oscuro
        'color_encabezado': '3CB371'  # Verde mar
    }
}

# --- Constante para altura de filas ---
ALTURA_FILA_DATOS = 25  # Misma altura que los encabezados

# --- Función para inicializar el archivo Excel con las tres tablas ---


def inicializar_excel():
    """Crea el archivo Excel con las tres tablas en una sola hoja"""
    if not os.path.exists(ARCHIVO_EXCEL):
        # Crear un libro nuevo
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = NOMBRE_HOJA

        fila_actual = 1

        for categoria, etiquetas in ETIQUETAS_CATEGORIA.items():
            # Título de la tabla
            ws.merge_cells(f'A{fila_actual}:D{fila_actual}')
            celda_titulo = ws[f'A{fila_actual}']
            celda_titulo.value = etiquetas['titulo']
            celda_titulo.font = Font(size=14, bold=True, color='FFFFFF')
            celda_titulo.fill = PatternFill(start_color=etiquetas['color_titulo'],
                                            end_color=etiquetas['color_titulo'],
                                            fill_type='solid')
            celda_titulo.alignment = Alignment(
                horizontal='center', vertical='center')
            ws.row_dimensions[fila_actual].height = 30

            fila_actual += 1

            # Encabezados de la tabla
            encabezados = ['Nombre', 'Cantidad', 'Fecha', 'Hora']
            for col, encabezado in enumerate(encabezados, start=1):
                celda = ws.cell(row=fila_actual, column=col)
                celda.value = encabezado
                celda.font = Font(bold=True, color='FFFFFF', size=11)
                celda.fill = PatternFill(start_color=etiquetas['color_encabezado'],
                                         end_color=etiquetas['color_encabezado'],
                                         fill_type='solid')
                celda.alignment = Alignment(
                    horizontal='center', vertical='center')
                celda.border = Border(
                    left=Side(style='thin', color='000000'),
                    right=Side(style='thin', color='000000'),
                    top=Side(style='thin', color='000000'),
                    bottom=Side(style='thin', color='000000')
                )

            ws.row_dimensions[fila_actual].height = ALTURA_FILA_DATOS
            fila_actual += 1

            # Agregar una fila vacía para datos (con altura definida)
            ws.row_dimensions[fila_actual].height = ALTURA_FILA_DATOS
            fila_actual += 1

            # Línea divisoria gruesa entre tablas
            if categoria != 'Dulces':
                for col in range(1, 5):
                    celda = ws.cell(row=fila_actual, column=col)
                    celda.border = Border(bottom=Side(
                        style='thick', color='000000'))
                    celda.value = ""
                ws.row_dimensions[fila_actual].height = 5
                fila_actual += 1

            # Espacio adicional entre tablas
            for _ in range(4):
                ws.row_dimensions[fila_actual].height = ALTURA_FILA_DATOS
                fila_actual += 1

        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 35  # Nombre
        ws.column_dimensions['B'].width = 12  # Cantidad
        ws.column_dimensions['C'].width = 15  # Fecha
        ws.column_dimensions['D'].width = 12  # Hora

        wb.save(ARCHIVO_EXCEL)
        return True
    return False

# --- Función CORREGIDA para guardar en la tabla correspondiente ---


def guardar_en_excel(categoria, nombre, cantidad, fecha=None, hora=None):
    """Añade un registro a la tabla específica con fecha y hora separados"""
    if fecha is None:
        fecha = datetime.now().strftime("%Y-%m-%d")
    if hora is None:
        hora = datetime.now().strftime("%H:%M:%S")

    try:
        # Asegurar que el archivo existe
        inicializar_excel()

        # Cargar el libro existente
        wb = load_workbook(ARCHIVO_EXCEL)
        ws = wb[NOMBRE_HOJA]

        # Encontrar la posición de la tabla de la categoría
        fila_tabla = encontrar_fila_tabla(
            ws, ETIQUETAS_CATEGORIA[categoria]['titulo'])

        if fila_tabla is None:
            return False, f"❌ No se encontró la tabla de {categoria}"

        # La tabla empieza 2 filas después del título (título + encabezados)
        fila_datos = fila_tabla + 2

        # Encontrar la primera fila vacía en esa tabla
        while True:
            celda = ws.cell(row=fila_datos, column=1)
            # Verificar si la celda está vacía Y no es parte de un merged cell
            if celda.value is None:
                # Verificar que no sea una celda combinada
                es_combinada = False
                for merged_range in ws.merged_cells.ranges:
                    if celda.coordinate in merged_range:
                        es_combinada = True
                        break

                if not es_combinada:
                    break

            fila_datos += 1

            # Límite de seguridad
            if fila_datos > 1000:
                return False, "❌ Error: No se encontró espacio para guardar"

        # *** IMPORTANTE: Establecer altura de la fila ANTES de insertar datos ***
        ws.row_dimensions[fila_datos].height = ALTURA_FILA_DATOS

        # Insertar el nuevo registro
        ws.cell(row=fila_datos, column=1, value=nombre)
        ws.cell(row=fila_datos, column=2, value=cantidad)
        ws.cell(row=fila_datos, column=3, value=fecha)
        ws.cell(row=fila_datos, column=4, value=hora)

        # Aplicar estilo a las celdas
        for col in range(1, 5):
            celda = ws.cell(row=fila_datos, column=col)
            celda.alignment = Alignment(
                horizontal='left', vertical='center', wrap_text=True)
            celda.border = Border(
                left=Side(style='thin', color='000000'),
                right=Side(style='thin', color='000000'),
                top=Side(style='thin', color='000000'),
                bottom=Side(style='thin', color='000000')
            )

        # Alternar colores de fondo para mejor legibilidad
        # Calculamos la posición relativa dentro de la tabla
        posicion_relativa = fila_datos - (fila_tabla + 2)
        if posicion_relativa % 2 == 0:
            for col in range(1, 5):
                ws.cell(row=fila_datos, column=col).fill = PatternFill(
                    start_color='F5F5F5', end_color='F5F5F5', fill_type='solid'
                )

        wb.save(ARCHIVO_EXCEL)
        return True, f"✅ Registro guardado en {categoria}"

    except Exception as e:
        return False, f"❌ Error al guardar: {str(e)}"

# --- Función para encontrar la fila donde empieza una tabla ---


def encontrar_fila_tabla(ws, titulo_buscado):
    """Encuentra la fila donde está el título de una tabla específica"""
    for row in range(1, ws.max_row + 1):
        celda = ws.cell(row=row, column=1)
        if celda.value == titulo_buscado:
            return row
    return None

# --- Función para filtrar registros por fecha y categoría ---


def filtrar_por_fecha(categoria, fecha_objetivo):
    """Obtiene los registros de una tabla específica filtrados por fecha"""
    if not os.path.exists(ARCHIVO_EXCEL):
        return []

    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        ws = wb[NOMBRE_HOJA]

        # Encontrar la tabla
        fila_tabla = encontrar_fila_tabla(
            ws, ETIQUETAS_CATEGORIA[categoria]['titulo'])

        if fila_tabla is None:
            wb.close()
            return []

        # Los datos empiezan 2 filas después del título
        fila_datos = fila_tabla + 2

        registros = []
        fecha_buscada = fecha_objetivo.strftime("%Y-%m-%d")

        # Leer fila por fila hasta encontrar una vacía o la siguiente tabla
        while True:
            nombre = ws.cell(row=fila_datos, column=1).value

            # Si encontramos un título de otra tabla, paramos
            es_titulo = False
            for cat, etq in ETIQUETAS_CATEGORIA.items():
                if nombre == etq['titulo']:
                    es_titulo = True
                    break

            if es_titulo or nombre is None:
                break

            cantidad = ws.cell(row=fila_datos, column=2).value
            fecha = ws.cell(row=fila_datos, column=3).value
            hora = ws.cell(row=fila_datos, column=4).value

            if fecha and fecha == fecha_buscada:
                registros.append(
                    [nombre, cantidad, fecha, hora if hora else ""])

            fila_datos += 1

        wb.close()

        # Ordenar por hora (más reciente primero)
        registros.sort(key=lambda x: x[3] if x[3] else "", reverse=True)
        return registros

    except Exception as e:
        print(f"Error filtrando: {e}")
        return []

# --- Función para crear ventana de calendario ---


def mostrar_calendario(fecha_actual):
    año_actual = fecha_actual.year
    mes_actual = fecha_actual.month

    cal = calendar.monthcalendar(año_actual, mes_actual)

    layout_calendario = [
        [sg.Text(f"{calendar.month_name[mes_actual]} {año_actual}",
                 font=('Arial', 16, 'bold'),
                 justification='center',
                 expand_x=True)],
        [sg.Text('')],

        [sg.Text(dia, size=(4, 1), justification='center', font=('Arial', 10, 'bold'))
         for dia in ['Lun', 'Mar', 'Mié', 'Jue', 'Vie', 'Sáb', 'Dom']],
    ]

    for semana in cal:
        fila_botones = []
        for dia in semana:
            if dia == 0:
                fila_botones.append(sg.Text('', size=(4, 2)))
            else:
                if (dia == fecha_actual.day and mes_actual == datetime.now().month
                        and año_actual == datetime.now().year):
                    color_boton = ('white', 'green')
                else:
                    color_boton = ('black', 'white')

                fila_botones.append(
                    sg.Button(str(dia),
                              size=(4, 2),
                              key=f'-DIA_CAL_{dia}-',
                              button_color=color_boton)
                )
        layout_calendario.append(fila_botones)

    layout_calendario.extend([
        [sg.Text('')],
        [sg.Button('◀ Mes Anterior', size=(12, 1), key='-MES_ANTERIOR-'),
         sg.Button('Hoy', size=(8, 1), key='-HOY-'),
         sg.Button('Mes Siguiente ▶', size=(12, 1), key='-MES_SIGUIENTE-')],
        [sg.Text('')],
        [sg.Button('Cancelar', size=(10, 1)),
         sg.Button('Seleccionar', size=(10, 1), button_color=('white', 'blue'))]
    ])

    window_cal = sg.Window('Seleccionar Fecha',
                           layout_calendario,
                           modal=True,
                           element_justification='center')

    fecha_seleccionada = None
    mes_temporal = mes_actual
    año_temporal = año_actual

    while True:
        event_cal, values_cal = window_cal.read()

        if event_cal in (sg.WINDOW_CLOSED, 'Cancelar'):
            break

        if event_cal == 'Seleccionar':
            if fecha_seleccionada:
                break
            else:
                sg.popup('Por favor selecciona un día primero', title='Aviso')

        if event_cal == '-HOY-':
            fecha_seleccionada = datetime.now()
            break

        if event_cal == '-MES_ANTERIOR-':
            mes_temporal -= 1
            if mes_temporal < 1:
                mes_temporal = 12
                año_temporal -= 1
            window_cal.close()
            fecha_temporal = datetime(año_temporal, mes_temporal, 1)
            fecha_seleccionada = mostrar_calendario(fecha_temporal)
            break

        if event_cal == '-MES_SIGUIENTE-':
            mes_temporal += 1
            if mes_temporal > 12:
                mes_temporal = 1
                año_temporal += 1
            window_cal.close()
            fecha_temporal = datetime(año_temporal, mes_temporal, 1)
            fecha_seleccionada = mostrar_calendario(fecha_temporal)
            break

        if event_cal.startswith('-DIA_CAL_'):
            dia = int(event_cal.replace('-DIA_CAL_', '').replace('-', ''))
            fecha_seleccionada = datetime(año_temporal, mes_temporal, dia)

    window_cal.close()
    return fecha_seleccionada


# --- Inicializar Excel al iniciar el programa ---
inicializar_excel()

# --- Obtener fecha y hora actual ---
fecha_actual = datetime.now()
fecha_texto = fecha_actual.strftime("%d de %B de %Y")
hora_actual = fecha_actual.strftime("%H:%M:%S")

# --- Variable global para categoría actual ---
categoria_actual = 'Vino'

# --- Diseño de la interfaz ---
layout = [
    # === SELECTOR DE CATEGORÍA ===
    [sg.Text('📦 SELECCIONAR CATEGORÍA', font=('Arial', 14, 'bold'),
             justification='center', expand_x=True)],

    [sg.Button('🍷 VINO', size=(15, 2), font=('Arial', 12, 'bold'),
               button_color=('white', 'darkred'), key='-CAT_VINO-'),
     sg.Button('🥫 LATAS', size=(15, 2), font=('Arial', 12, 'bold'),
               button_color=('white', 'darkblue'), key='-CAT_LATAS-'),
     sg.Button('🍬 DULCES', size=(15, 2), font=('Arial', 12, 'bold'),
               button_color=('white', 'darkgreen'), key='-CAT_DULCES-')],

    [sg.HorizontalSeparator()],

    # === SELECTOR DE FECHA ===
    [sg.Text('📅 FILTRAR POR FECHA', font=('Arial', 12, 'bold'),
             justification='center', expand_x=True)],

    [sg.Button('📆 Seleccionar Fecha', size=(20, 2), font=('Arial', 11, 'bold'),
               button_color=('white', 'purple'), key='-SELECCIONAR_FECHA-'),
     sg.Text(fecha_texto, key='-FECHA_MOSTRADA-', font=('Arial', 14, 'bold'),
             size=(30, 1), justification='center')],

    [sg.HorizontalSeparator()],
    [sg.Text('')],

    # === CONTENIDO PRINCIPAL ===
    [
        # COLUMNA IZQUIERDA: Tabla de registros
        sg.Column([
            [sg.Text('📋 REGISTROS DEL DÍA', font=('Arial', 12, 'bold'),
                     key='-TITULO_TABLA-')],
            [sg.Table(
                values=[],
                headings=['Nombre', 'Cantidad', 'Fecha', 'Hora'],
                max_col_width=30,
                auto_size_columns=False,
                col_widths=[30, 12, 15, 10],
                display_row_numbers=True,
                justification='left',
                num_rows=20,
                key='-TABLA-',
                row_height=25,
                enable_events=True,
                expand_x=True,
                expand_y=True
            )],
            [sg.Button('🔄 Actualizar', size=(15, 1)),
             sg.Button('📂 Abrir Excel', size=(15, 1),
                       button_color=('white', 'green')),
             sg.Button('📊 Ver Todos', size=(15, 1), button_color=('white', 'orange'))]
        ], expand_x=True, expand_y=True, scrollable=True),

        # COLUMNA DERECHA: Formulario con etiquetas dinámicas
        sg.Column([
            [sg.Text(f'✨ NUEVO REGISTRO - {categoria_actual.upper()}',
                     font=('Arial', 12, 'bold'), key='-TITULO_FORMULARIO-')],
            [sg.Text('')],

            [sg.Text(ETIQUETAS_CATEGORIA[categoria_actual]['nombre'],
                     size=(20, 1), key='-LABEL_NOMBRE-', font=('Arial', 11))],
            [sg.InputText(key='-NOMBRE-', size=(30, 1), font=('Arial', 11))],

            [sg.Text('')],

            [sg.Text(ETIQUETAS_CATEGORIA[categoria_actual]['cantidad'],
                     size=(20, 1), key='-LABEL_CANTIDAD-', font=('Arial', 11))],
            [sg.InputText(key='-CANTIDAD-', size=(30, 1), font=('Arial', 11))],

            [sg.Text('')],

            # CAMPOS SEPARADOS DE FECHA Y HORA
            [sg.Text('Fecha de registro:', size=(20, 1))],
            [sg.InputText(key='-FECHA_REGISTRO-', size=(30, 1),
                          default_text=fecha_actual.strftime("%Y-%m-%d"),
                          font=('Arial', 10))],

            [sg.Text('')],

            [sg.Text('Hora de registro:', size=(20, 1))],
            [sg.InputText(key='-HORA_REGISTRO-', size=(30, 1),
                          default_text=hora_actual,
                          font=('Arial', 10))],

            [sg.Text('')],
            [sg.Text('')],

            [sg.Button('💾 Guardar Registro', size=(20, 2),
                       button_color=('white', 'blue'),
                       font=('Arial', 11, 'bold'),
                       bind_return_key=True,
                       key='-GUARDAR-')],

            [sg.Text('')],

            [sg.Text('', key='-MENSAJE-', size=(40, 2), text_color='green',
                     font=('Arial', 10), justification='center')],

            [sg.Text('')],
            [sg.Button('🗑️ Limpiar campos', size=(20, 1),
                       button_color=('white', 'gray'))],
            [sg.Button('🕐 Usar hora actual', size=(20, 1),
                       button_color=('white', 'blue'), key='-HORA_ACTUAL-')]
        ], element_justification='center', vertical_alignment='top')
    ],

    # === BARRA DE ESTADO ===
    [sg.HorizontalSeparator()],
    [sg.Text('💡 Tip: Todas las filas tienen ahora la misma altura para mejor presentación',
             text_color='gray', font=('Arial', 9), expand_x=True)]
]

# --- Crear ventana principal ---
window = sg.Window('Tienda Registro - Sistema con Fecha y Hora Separados',
                   layout,
                   size=(1100, 680),
                   resizable=True,
                   finalize=True)

# --- Variables globales ---
fecha_seleccionada_global = fecha_actual

# --- Función para actualizar la tabla con registros filtrados ---


def actualizar_tabla_por_fecha(categoria, fecha_filtro):
    registros = filtrar_por_fecha(categoria, fecha_filtro)
    window['-TABLA-'].update(values=registros)

    # Actualizar título de la tabla
    fecha_formateada = fecha_filtro.strftime("%d/%m/%Y")
    window['-TITULO_TABLA-'].update(
        f'📋 {categoria.upper()} - Registros del {fecha_formateada}')

    return len(registros) > 0

# --- Función para ver todos los registros de una categoría ---


def ver_todos_registros(categoria):
    if not os.path.exists(ARCHIVO_EXCEL):
        return False

    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        ws = wb[NOMBRE_HOJA]

        fila_tabla = encontrar_fila_tabla(
            ws, ETIQUETAS_CATEGORIA[categoria]['titulo'])

        if fila_tabla is None:
            wb.close()
            return False

        fila_datos = fila_tabla + 2
        registros = []

        while True:
            nombre = ws.cell(row=fila_datos, column=1).value

            # Si encontramos un título de otra tabla, paramos
            es_titulo = False
            for cat, etq in ETIQUETAS_CATEGORIA.items():
                if nombre == etq['titulo']:
                    es_titulo = True
                    break

            if es_titulo or nombre is None:
                break

            cantidad = ws.cell(row=fila_datos, column=2).value
            fecha = ws.cell(row=fila_datos, column=3).value
            hora = ws.cell(row=fila_datos, column=4).value
            registros.append([nombre, cantidad, fecha, hora if hora else ""])
            fila_datos += 1

        wb.close()

        if registros:
            registros.sort(key=lambda x: f"{x[2]} {x[3]}", reverse=True)
            window['-TABLA-'].update(values=registros)
            window['-TITULO_TABLA-'].update(
                f'📋 {categoria.upper()} - TODOS LOS REGISTROS')
            return True

    except Exception as e:
        sg.popup_error(f"Error al cargar: {e}")
    return False

# --- Función para cambiar de categoría ---


def cambiar_categoria(nueva_categoria):
    global categoria_actual
    categoria_actual = nueva_categoria

    # Actualizar etiquetas del formulario
    window['-LABEL_NOMBRE-'].update(
        ETIQUETAS_CATEGORIA[categoria_actual]['nombre'])
    window['-LABEL_CANTIDAD-'].update(
        ETIQUETAS_CATEGORIA[categoria_actual]['cantidad'])

    # Actualizar título del formulario
    window['-TITULO_FORMULARIO-'].update(
        f'✨ NUEVO REGISTRO - {categoria_actual.upper()}')

    # Actualizar tabla con la nueva categoría
    actualizar_tabla_por_fecha(categoria_actual, fecha_seleccionada_global)

    # Mensaje de confirmación
    window['-MENSAJE-'].update(
        f'✅ Categoría cambiada a {categoria_actual}', text_color='green')


# --- Cargar registros iniciales ---
actualizar_tabla_por_fecha(categoria_actual, fecha_actual)

# --- Bucle principal ---
while True:
    event, values = window.read()

    if event == sg.WINDOW_CLOSED:
        break

    # Actualizar fecha mostrada
    window['-FECHA_MOSTRADA-'].update(
        fecha_seleccionada_global.strftime("%d de %B de %Y"))

    # Cambiar categoría
    if event == '-CAT_VINO-':
        cambiar_categoria('Vino')
    elif event == '-CAT_LATAS-':
        cambiar_categoria('Latas')
    elif event == '-CAT_DULCES-':
        cambiar_categoria('Dulces')

    # Seleccionar fecha del calendario
    if event == '-SELECCIONAR_FECHA-':
        fecha_elegida = mostrar_calendario(fecha_seleccionada_global)
        if fecha_elegida:
            fecha_seleccionada_global = fecha_elegida
            window['-FECHA_REGISTRO-'].update(
                fecha_elegida.strftime("%Y-%m-%d"))
            actualizar_tabla_por_fecha(
                categoria_actual, fecha_seleccionada_global)
            window['-MENSAJE-'].update(f"✅ Mostrando registros del {fecha_elegida.strftime('%d/%m/%Y')}",
                                       text_color='green')

    # Usar hora actual
    if event == '-HORA_ACTUAL-':
        hora_actual = datetime.now().strftime("%H:%M:%S")
        window['-HORA_REGISTRO-'].update(hora_actual)
        window['-MENSAJE-'].update("🕐 Hora actual actualizada",
                                   text_color='blue')

    # Guardar registro
    if event == '-GUARDAR-':
        nombre = values['-NOMBRE-'].strip()
        cantidad = values['-CANTIDAD-'].strip()
        fecha = values['-FECHA_REGISTRO-'].strip()
        hora = values['-HORA_REGISTRO-'].strip()

        if not nombre or not cantidad or not fecha or not hora:
            window['-MENSAJE-'].update(
                "❌ Todos los campos son obligatorios", text_color='red')
        elif not cantidad.isdigit():
            window['-MENSAJE-'].update("❌ La cantidad debe ser un número",
                                       text_color='red')
        else:
            exito, mensaje = guardar_en_excel(
                categoria_actual, nombre, int(cantidad), fecha, hora)

            if exito:
                window['-MENSAJE-'].update(mensaje, text_color='green')
                window['-NOMBRE-'].update('')
                window['-CANTIDAD-'].update('')
                actualizar_tabla_por_fecha(
                    categoria_actual, fecha_seleccionada_global)
                window['-NOMBRE-'].set_focus()
            else:
                window['-MENSAJE-'].update(mensaje, text_color='red')

    # Actualizar tabla
    if event == '🔄 Actualizar':
        actualizar_tabla_por_fecha(categoria_actual, fecha_seleccionada_global)
        window['-MENSAJE-'].update("✅ Tabla actualizada", text_color='green')

    # Ver todos los registros
    if event == '📊 Ver Todos':
        if ver_todos_registros(categoria_actual):
            window['-MENSAJE-'].update(f"✅ Mostrando todos los registros de {categoria_actual}",
                                       text_color='green')
        else:
            window['-MENSAJE-'].update(f"ℹ️ No hay registros en {categoria_actual}",
                                       text_color='orange')

    # Abrir Excel
    if event == '📂 Abrir Excel':
        if os.path.exists(ARCHIVO_EXCEL):
            os.startfile(ARCHIVO_EXCEL)
        else:
            sg.popup("📁 Aún no hay archivo Excel. Guarda tu primer registro.",
                     title="Información")

    # Limpiar campos
    if event == '🗑️ Limpiar campos':
        window['-NOMBRE-'].update('')
        window['-CANTIDAD-'].update('')
        window['-MENSAJE-'].update("🧹 Campos limpiados", text_color='blue')
        window['-NOMBRE-'].set_focus()

window.close()
