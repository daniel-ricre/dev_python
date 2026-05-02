import PySimpleGUI as sg
import pandas as pd
import os
from datetime import datetime
import calendar
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- TEMA PROFESIONAL OSCURO (Estilo Linear/Notion) ---
sg.theme('DarkGrey5')  # Base oscura elegante
sg.set_options(font=('Inter', 10), dpi_awareness=True)

# --- Configuración Excel ---
ARCHIVO_EXCEL = 'registro_inventario.xlsx'
HOJAS_EXCEL = {'Vino': 'VINO', 'Latas': 'LATAS', 'Dulces': 'DULCES'}

CATEGORIAS = {
    'Vino': {'nombre': 'Vino', 'cantidad': 'Botellas', 'icono': '🍷', 'color': '#E11D48'},
    'Latas': {'nombre': 'Lata', 'cantidad': 'Unidades', 'icono': '🥫', 'color': '#2563EB'},
    'Dulces': {'nombre': 'Dulce', 'cantidad': 'Piezas', 'icono': '🍬', 'color': '#059669'}
}

# --- Funciones Excel (Versión Robusta) ---


def inicializar_excel():
    if not os.path.exists(ARCHIVO_EXCEL):
        wb = Workbook()
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        for cat, hoja in HOJAS_EXCEL.items():
            ws = wb.create_sheet(hoja)
            ws['A1'] = 'Nombre'
            ws['B1'] = 'Cantidad'
            ws['C1'] = 'Fecha'
            ws['D1'] = 'Hora'
            # Estilos básicos para que se vea bien
            for col in range(1, 5):
                cell = ws.cell(row=1, column=col)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill(
                    start_color='1F2937', end_color='1F2937', fill_type='solid')
                cell.alignment = Alignment(horizontal='center')
            ws.column_dimensions['A'].width = 30
            ws.column_dimensions['B'].width = 12
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 12
        wb.save(ARCHIVO_EXCEL)


def guardar_registro(categoria, nombre, cantidad, fecha_str):
    try:
        inicializar_excel()
        wb = load_workbook(ARCHIVO_EXCEL)
        ws = wb[HOJAS_EXCEL[categoria]]
        fila = ws.max_row + 1
        hora = datetime.now().strftime("%H:%M")
        ws.cell(row=fila, column=1, value=nombre)
        ws.cell(row=fila, column=2, value=cantidad)
        ws.cell(row=fila, column=3, value=fecha_str)
        ws.cell(row=fila, column=4, value=hora)
        wb.save(ARCHIVO_EXCEL)
        return True, "Guardado"
    except Exception as e:
        return False, str(e)


def leer_registros(categoria, fecha_filtro=None):
    if not os.path.exists(ARCHIVO_EXCEL):
        return []
    try:
        df = pd.read_excel(ARCHIVO_EXCEL, sheet_name=HOJAS_EXCEL[categoria])
        if df.empty:
            return []
        if fecha_filtro:
            df = df[df['Fecha'] == fecha_filtro]
        # Ordenar por hora descendente
        df = df.sort_values(by=['Fecha', 'Hora'], ascending=[False, False])
        return df.values.tolist()
    except:
        return []

# --- Calendario Minimalista (Panel Lateral) ---


def crear_calendario_layout(fecha_actual, fecha_sel):
    año, mes = fecha_actual.year, fecha_actual.month
    cal = calendar.monthcalendar(año, mes)

    layout = [
        [sg.Text(f"{calendar.month_name[mes]} {año}", font=(
            'Inter', 12, 'bold'), text_color='#E2E8F0')],
        [sg.Text("L  M  X  J  V  S  D", font=(
            'Inter', 8), text_color='#64748B')]
    ]

    for semana in cal:
        fila = []
        for dia in semana:
            if dia == 0:
                fila.append(sg.Text('  ', size=(2, 1)))
            else:
                fecha_boton = datetime(año, mes, dia)
                # Lógica de colores
                if fecha_boton.date() == fecha_sel.date():
                    color_btn = (CATEGORIAS['Vino']
                                 ['color'], '#FFFFFF')  # Invertido
                    txt_color = '#FFFFFF'
                elif fecha_boton.date() == datetime.now().date():
                    color_btn = ('#334155', '#38BDF8')
                    txt_color = '#38BDF8'
                else:
                    color_btn = ('#1E293B', '#64748B')
                    txt_color = '#64748B'

                btn = sg.Button(str(dia), size=(2, 1), key=f'-CAL_{dia}-',
                                button_color=color_btn, border_width=0, font=('Inter', 9, 'bold'),
                                pad=(0, 2))
                fila.append(btn)
        layout.append(fila)

    layout.append([sg.Button('Hoy', size=(8, 1), key='-HOY-', button_color=('#1E293B', '#475569'), border_width=0),
                   sg.Button('<', size=(3, 1), key='-MES_ANT-'), sg.Button('>', size=(3, 1), key='-MES_SIG-')])
    return layout


# --- Construcción de la Interfaz Principal ---
inicializar_excel()
fecha_global = datetime.now()
fecha_seleccionada = fecha_global
categoria_activa = 'Vino'

# Colores dinámicos
color_activo = CATEGORIAS[categoria_activa]['color']

# === BARRA LATERAL IZQUIERDA (Sidebar) ===
sidebar = [
    [sg.Text('INVENTORY', font=('Inter', 14, 'bold'), text_color='#94A3B8')],
    [sg.Text('')],

    # Botones de Categoría (Estilo Menú Vertical)
    [sg.Button(f"{CATEGORIAS['Vino']['icono']}  Vinos", size=(14, 2), key='-CAT_VINO-',
               button_color=('#FFFFFF' if categoria_activa == 'Vino' else '#94A3B8',
                             CATEGORIAS['Vino']['color'] if categoria_activa == 'Vino' else '#1E293B'),
               border_width=0, font=('Inter', 11))],
    [sg.Button(f"{CATEGORIAS['Latas']['icono']}  Latas", size=(14, 2), key='-CAT_LATAS-',
               button_color=('#FFFFFF' if categoria_activa == 'Latas' else '#94A3B8',
                             CATEGORIAS['Latas']['color'] if categoria_activa == 'Latas' else '#1E293B'),
               border_width=0, font=('Inter', 11))],
    [sg.Button(f"{CATEGORIAS['Dulces']['icono']}  Dulces", size=(14, 2), key='-CAT_DULCES-',
               button_color=('#FFFFFF' if categoria_activa == 'Dulces' else '#94A3B8',
                             CATEGORIAS['Dulces']['color'] if categoria_activa == 'Dulces' else '#1E293B'),
               border_width=0, font=('Inter', 11))],

    [sg.Text('_' * 20, text_color='#334155')],
    [sg.Text('📅 CALENDARIO', font=('Inter', 10, 'bold'), text_color='#94A3B8')],
]

# === PANEL CENTRAL (Tabla) ===
tabla_panel = [
    [sg.Text(f'📋 Registros de {categoria_activa}', font=(
        'Inter', 13, 'bold'), key='-TITULO_TABLA-', text_color='#F1F5F9')],
    [sg.Text(f'Mostrando datos del: {fecha_seleccionada.strftime("%d/%m/%Y")}',
             key='-FECHA_TITULO-', text_color='#64748B')],
    [sg.Table(
        values=leer_registros(
            categoria_activa, fecha_seleccionada.strftime("%Y-%m-%d")),
        headings=['Producto', 'Cant', 'Fecha', 'Hora'],
        col_widths=[30, 10, 15, 10],
        auto_size_columns=False,
        display_row_numbers=True,
        num_rows=18,
        key='-TABLA-',
        row_height=30,
        font=('Inter', 10),
        header_font=('Inter', 9, 'bold'),
        header_text_color='#94A3B8',
        header_background_color='#1E293B',
        alternating_row_color='#1E293B',
        selected_row_colors=('#FFFFFF', CATEGORIAS[categoria_activa]['color']),
        expand_x=True, expand_y=True,
        enable_events=True
    )],
    [sg.Button('🔄 Actualizar', size=(12, 1), button_color=('#1E293B', '#334155'), border_width=0),
     sg.Button('📂 Abrir Excel', size=(12, 1), button_color=(
         '#1E293B', '#334155'), border_width=0),
     sg.Push(),  # Empuja los elementos a la derecha
     sg.Button('📊 Ver Todos', size=(12, 1), button_color=('#1E293B', '#334155'), border_width=0, key='-VER_TODOS-')]
]

# === PANEL DERECHO (Formulario) ===
form_panel = [
    [sg.Text(f'Nuevo Registro', font=('Inter', 13, 'bold'), text_color='#F1F5F9')],
    [sg.Text(f'Categoría: {categoria_activa}', key='-CAT_FORM-',
             text_color=CATEGORIAS[categoria_activa]['color'])],
    [sg.Text('_' * 25, text_color='#334155')],

    [sg.Text(f"Nombre del {CATEGORIAS[categoria_activa]['nombre']}", font=(
        'Inter', 10), text_color='#94A3B8')],
    [sg.Input(key='-NOMBRE-', size=(25, 1), font=('Inter', 11),
              background_color='#1E293B', text_color='#F1F5F9', border_width=0)],

    [sg.Text(f"Cantidad ({CATEGORIAS[categoria_activa]['cantidad']})", font=(
        'Inter', 10), text_color='#94A3B8')],
    [sg.Input(key='-CANTIDAD-', size=(25, 1), font=('Inter', 11),
              background_color='#1E293B', text_color='#F1F5F9', border_width=0)],

    [sg.Text('Fecha', font=('Inter', 10), text_color='#94A3B8')],
    [sg.Input(key='-FECHA-', default_text=fecha_seleccionada.strftime("%Y-%m-%d"), size=(25, 1),
              font=('Inter', 11), background_color='#1E293B', text_color='#F1F5F9', border_width=0)],

    [sg.Text('')],
    [sg.Button('💾  GUARDAR PRODUCTO', size=(20, 2), key='-GUARDAR-',
               button_color=('#FFFFFF', CATEGORIAS[categoria_activa]['color']),
               border_width=0, font=('Inter', 11, 'bold'))],

    [sg.Text('', key='-MSG-', size=(30, 2),
             text_color='#10B981', font=('Inter', 9))]
]

# === ENSAMBLAJE FINAL ===
layout = [
    [sg.Column(sidebar, vertical_alignment='top', pad=(20, 20), background_color='#0F172A'),
     sg.VSeparator(color='#334155'),
     sg.Column(tabla_panel, expand_x=True, expand_y=True,
               pad=(20, 20), background_color='#0F172A'),
     sg.VSeparator(color='#334155'),
     sg.Column(form_panel, vertical_alignment='top', pad=(20, 20), background_color='#0F172A', element_justification='center')]
]

window = sg.Window('Inventory Pro', layout, size=(1300, 700),
                   margins=(0, 0), background_color='#0F172A',
                   resizable=True, finalize=True)

# --- Función de refresco de UI Completa (Cambia colores dinámicos) ---


def refrescar_ui():
    global color_activo
    color_activo = CATEGORIAS[categoria_activa]['color']

    # Actualizar textos y colores
    window['-TITULO_TABLA-'].update(f'📋 Registros de {categoria_activa}')
    window['-CAT_FORM-'].update(
        f'Categoría: {categoria_activa}', text_color=color_activo)

    # Actualizar botones de menú lateral (quitar color al anterior)
    for cat in ['VINO', 'LATAS', 'DULCES']:
        is_active = (cat == categoria_activa.upper())
        color_btn = color_activo if is_active else '#1E293B'
        txt_color = '#FFFFFF' if is_active else '#94A3B8'
        window[f'-CAT_{cat}-'].update(button_color=(txt_color, color_btn))

    # Actualizar botón Guardar
    window['-GUARDAR-'].update(button_color=('#FFFFFF', color_activo))
    # Actualizar color de selección de tabla
    window['-TABLA-'].update(selected_row_colors=('#FFFFFF', color_activo))

    # Recargar datos
    datos = leer_registros(
        categoria_activa, fecha_seleccionada.strftime("%Y-%m-%d"))
    window['-TABLA-'].update(values=datos)


# --- BUCLE PRINCIPAL ---
while True:
    event, values = window.read()

    if event == sg.WINDOW_CLOSED:
        break

    # Navegación de Categorías
    if event == '-CAT_VINO-':
        categoria_activa = 'Vino'
        refrescar_ui()
    if event == '-CAT_LATAS-':
        categoria_activa = 'Latas'
        refrescar_ui()
    if event == '-CAT_DULCES-':
        categoria_activa = 'Dulces'
        refrescar_ui()

    # Actualizar Fecha desde Calendario
    if event.startswith('-CAL_'):
        dia = int(event.replace('-CAL_', '').replace('-', ''))
        fecha_seleccionada = datetime(
            fecha_global.year, fecha_global.month, dia)
        window['-FECHA-'].update(fecha_seleccionada.strftime("%Y-%m-%d"))
        window['-FECHA_TITULO-'].update(
            f'Mostrando datos del: {fecha_seleccionada.strftime("%d/%m/%Y")}')
        datos = leer_registros(
            categoria_activa, fecha_seleccionada.strftime("%Y-%m-%d"))
        window['-TABLA-'].update(values=datos)

    if event == '-HOY-':
        fecha_seleccionada = datetime.now()
        fecha_global = fecha_seleccionada
        window['-FECHA-'].update(fecha_seleccionada.strftime("%Y-%m-%d"))
        # Aquí normalmente recargarías el calendario, pero lo omitimos para brevedad

    # Guardar Producto
    if event == '-GUARDAR-':
        nombre = values['-NOMBRE-'].strip()
        cantidad = values['-CANTIDAD-'].strip()
        fecha = values['-FECHA-'].strip()

        if not nombre or not cantidad:
            window['-MSG-'].update('❌ Completa todos los campos',
                                   text_color='#EF4444')
        elif not cantidad.isdigit():
            window['-MSG-'].update('❌ Cantidad inválida', text_color='#EF4444')
        else:
            exito, msg = guardar_registro(
                categoria_activa, nombre, int(cantidad), fecha)
            if exito:
                window['-MSG-'].update(f'✅ {nombre} guardado',
                                       text_color='#10B981')
                window['-NOMBRE-'].update('')
                window['-CANTIDAD-'].update('')
                # Refrescar Tabla
                datos = leer_registros(
                    categoria_activa, fecha_seleccionada.strftime("%Y-%m-%d"))
                window['-TABLA-'].update(values=datos)
                window['-NOMBRE-'].set_focus()
            else:
                window['-MSG-'].update(f'❌ Error: {msg}', text_color='#EF4444')

    # Ver Todos
    if event == '-VER_TODOS-':
        datos = leer_registros(categoria_activa)
        window['-TABLA-'].update(values=datos)
        window['-FECHA_TITULO-'].update('Mostrando TODOS los registros')

    # Abrir Excel
    if event == '📂 Abrir Excel':
        if os.path.exists(ARCHIVO_EXCEL):
            os.startfile(ARCHIVO_EXCEL)

    if event == '🔄 Actualizar':
        datos = leer_registros(
            categoria_activa, fecha_seleccionada.strftime("%Y-%m-%d"))
        window['-TABLA-'].update(values=datos)
        window['-MSG-'].update('✅ Datos actualizados', text_color='#10B981')

window.close()
