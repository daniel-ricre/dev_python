import PySimpleGUI as sg
import pandas as pd
import os
from datetime import datetime
import calendar
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ------------------------------------------------------------
# TEMA Y CONSTANTES DE DISEÑO (PROFESIONAL)
# ------------------------------------------------------------
sg.theme('DarkBlack1')
sg.set_options(font=('Segoe UI', 10), dpi_awareness=True)

COLORS = {
    'bg': '#0A0A0A',
    'surface': '#18181B',
    'border': '#27272A',
    'text': '#FAFAFA',
    'text_secondary': '#A1A1AA',
    'accent': '#3B82F6',
    'success': '#10B981',
    'warning': '#F59E0B',
    'danger': '#EF4444'
}

CATEGORIAS = {
    'Vino':   {'nombre': 'Vino',   'medida': 'Botellas', 'icono': '🍷', 'color': '#E11D48'},
    'Latas':  {'nombre': 'Lata',   'medida': 'Unidades', 'icono': '🥫', 'color': '#2563EB'},
    'Dulces': {'nombre': 'Dulce',  'medida': 'Piezas',   'icono': '🍬', 'color': '#059669'}
}

ARCHIVO_EXCEL = 'inventario_db.xlsx'
HOJAS_EXCEL = {'Vino': 'VINO', 'Latas': 'LATAS', 'Dulces': 'DULCES'}

# ------------------------------------------------------------
# FUNCIONES DE EXCEL (ROBUSTAS)
# ------------------------------------------------------------


def init_excel():
    if not os.path.exists(ARCHIVO_EXCEL):
        with pd.ExcelWriter(ARCHIVO_EXCEL, engine='openpyxl') as writer:
            for hoja in HOJAS_EXCEL.values():
                pd.DataFrame(columns=['Producto', 'Cantidad', 'Fecha', 'Hora']).to_excel(
                    writer, sheet_name=hoja, index=False
                )
        # Estilizar el archivo
        wb = load_workbook(ARCHIVO_EXCEL)
        for hoja in HOJAS_EXCEL.values():
            ws = wb[hoja]
            ws.column_dimensions['A'].width = 35
            ws.column_dimensions['B'].width = 12
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 12
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill(
                    start_color='27272A', end_color='27272A', fill_type='solid')
                cell.alignment = Alignment(horizontal='center')
        wb.save(ARCHIVO_EXCEL)


def guardar_registro(categoria, nombre, cantidad):
    try:
        fecha = datetime.now().strftime("%Y-%m-%d")
        hora = datetime.now().strftime("%H:%M")
        df_nuevo = pd.DataFrame([[nombre, cantidad, fecha, hora]],
                                columns=['Producto', 'Cantidad', 'Fecha', 'Hora'])
        with pd.ExcelWriter(ARCHIVO_EXCEL, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
            writer.workbook = load_workbook(ARCHIVO_EXCEL)
            hoja = HOJAS_EXCEL[categoria]
            df_nuevo.to_excel(writer, sheet_name=hoja, index=False, header=False,
                              startrow=writer.sheets[hoja].max_row)
        return True, "Guardado correctamente"
    except Exception as e:
        return False, str(e)


def obtener_registros(categoria, fecha_filtro=None):
    if not os.path.exists(ARCHIVO_EXCEL):
        return []
    try:
        df = pd.read_excel(ARCHIVO_EXCEL, sheet_name=HOJAS_EXCEL[categoria])
        if df.empty:
            return []
        if fecha_filtro:
            df = df[df['Fecha'] == fecha_filtro]
        return df.sort_values(by=['Fecha', 'Hora'], ascending=[False, False]).values.tolist()
    except Exception:
        return []


# ------------------------------------------------------------
# INICIALIZACIÓN DE DATOS
# ------------------------------------------------------------
init_excel()
fecha_actual = datetime.now()
fecha_seleccionada = fecha_actual
categoria_activa = 'Vino'

# ------------------------------------------------------------
# CONSTRUCCIÓN DE LA INTERFAZ GRÁFICA (ESTÁTICA Y SÓLIDA)
# ------------------------------------------------------------

# --- Barra lateral: categorías ---


def crear_botones_categoria():
    btns = []
    for cat_key, cat_val in CATEGORIAS.items():
        is_active = (cat_key == categoria_activa)
        bg = cat_val['color'] if is_active else '#18181B'
        txt = '#FFFFFF' if is_active else COLORS['text_secondary']
        btns.append([sg.Button(f"{cat_val['icono']}  {cat_key.upper()}",
                               size=(15, 2),
                               key=f'-CAT-{cat_key.upper()}',
                               button_color=(txt, bg),
                               border_width=0,
                               font=('Segoe UI', 11, 'bold'))])
    return btns


# Contenedor de botones de categoría (se actualizarán colores, no los botones)
btns_categoria = crear_botones_categoria()

# --- Calendario (se crearán botones estáticos y luego se actualizarán) ---
# Calculamos el mes inicial
cal = calendar.monthcalendar(fecha_actual.year, fecha_actual.month)

# Construir encabezados del calendario
cal_header = [
    [sg.Text(f"{calendar.month_name[fecha_actual.month].upper()} {fecha_actual.year}",
             font=('Segoe UI', 13, 'bold'), text_color=COLORS['text'], key='-CAL_TITLE-')],
    [sg.Text("L  M  X  J  V  S  D", font=('Segoe UI', 8),
             text_color=COLORS['text_secondary'], pad=(5, 5))]
]

# Crear botones del calendario (hasta 6 semanas, 7 días por semana)
MAX_SEMANAS = 6
botones_calendario = []
for i in range(MAX_SEMANAS):
    fila_botones = []
    for j in range(7):
        # Para cada celda, un botón con texto vacío inicial si no hay día
        dia_val = 0
        if i < len(cal) and j < len(cal[i]):
            dia_val = cal[i][j]

        txt_dia = str(dia_val) if dia_val != 0 else ''
        btn = sg.Button(txt_dia, size=(2, 1), key=f'-CALBTN-{i}-{j}',
                        button_color=(COLORS['text'], '#18181B'), border_width=0,
                        font=('Segoe UI', 9, 'bold'), pad=(1, 2), disabled=(dia_val == 0))
        fila_botones.append(btn)
    botones_calendario.append(fila_botones)

# Navegación del calendario
cal_nav = [sg.Button('◀', size=(2, 1), key='-MES_ANT-', border_width=0, pad=(0, 5)),
           sg.Button('▶', size=(2, 1), key='-MES_SIG-', border_width=0),
           sg.Push(),
           sg.Button('HOY', size=(5, 1), key='-HOY-', border_width=0,
                     button_color=(COLORS['text'], '#27272A'))]

# --- Panel central: tabla de datos ---
tabla_panel = [
    [sg.Text(f'📋 INVENTARIO DE {categoria_activa.upper()}', font=('Segoe UI', 14, 'bold'),
             text_color=COLORS['text'], key='-TITLE-')],
    [sg.Text(f'Vista del: {fecha_seleccionada.strftime("%d/%m/%Y")}', font=('Segoe UI', 10),
             text_color=COLORS['text_secondary'], key='-SUBTITLE-')],
    [sg.Table(
        values=obtener_registros(
            categoria_activa, fecha_seleccionada.strftime("%Y-%m-%d")),
        headings=['Producto', 'Cant', 'Fecha', 'Hora'],
        col_widths=[35, 10, 15, 10],
        auto_size_columns=False,
        num_rows=20,
        key='-TABLE-',
        row_height=30,
        font=('Segoe UI', 10),
        header_font=('Segoe UI', 9, 'bold'),
        header_text_color='#A1A1AA',
        header_background_color='#18181B',
        alternating_row_color='#18181B',
        selected_row_colors=('#FFFFFF', CATEGORIAS[categoria_activa]['color']),
        expand_x=True, expand_y=True
    )],
    [sg.Button('🔄 Actualizar', size=(12, 1), button_color=(COLORS['text'], '#27272A'), border_width=0),
     sg.Button('📂 Abrir Excel', size=(12, 1), button_color=(
         COLORS['text'], '#27272A'), border_width=0),
     sg.Push(),
     sg.Button('📊 Ver Todos', size=(12, 1), button_color=(COLORS['text'], '#27272A'), border_width=0, key='-VER_TODOS-')]
]

# --- Panel derecho: formulario ---
form_panel = [
    [sg.Text('AÑADIR PRODUCTO', font=('Segoe UI', 13, 'bold'),
             text_color=COLORS['text'])],
    [sg.Text(f'Categoría activa: {categoria_activa}', font=('Segoe UI', 10),
             text_color=CATEGORIAS[categoria_activa]['color'], key='-CAT_LABEL-')],
    [sg.HorizontalSeparator(color=COLORS['border'])],
    [sg.Text(f"Nombre del {CATEGORIAS[categoria_activa]['nombre']}",
             text_color=COLORS['text_secondary'])],
    [sg.Input(key='-NOMBRE-', size=(25, 1), background_color='#18181B',
              text_color=COLORS['text'], border_width=0)],
    [sg.Text(f"Cantidad ({CATEGORIAS[categoria_activa]['medida']})",
             text_color=COLORS['text_secondary'])],
    [sg.Input(key='-CANTIDAD-', size=(25, 1), background_color='#18181B',
              text_color=COLORS['text'], border_width=0)],
    [sg.Text('Fecha', text_color=COLORS['text_secondary'])],
    [sg.Input(key='-FECHA-', default_text=fecha_seleccionada.strftime("%Y-%m-%d"), size=(25, 1),
              background_color='#18181B', text_color=COLORS['text'], border_width=0)],
    [sg.Text('')],
    [sg.Button('💾  GUARDAR', size=(20, 2), key='-GUARDAR-',
               button_color=('#FFFFFF', CATEGORIAS[categoria_activa]['color']), border_width=0,
               font=('Segoe UI', 11, 'bold'))],
    [sg.Text('', key='-MSG-', size=(30, 2),
             text_color=COLORS['success'], font=('Segoe UI', 9))]
]

# --- Columna lateral completa ---
sidebar_col = sg.Column([
    [sg.Text('MENÚ', font=('Segoe UI', 10, 'bold'),
             text_color=COLORS['text_secondary'])],
    *btns_categoria,
    [sg.Text('_' * 20, text_color=COLORS['border'])],
    [sg.Text('CALENDARIO', font=('Segoe UI', 10, 'bold'),
             text_color=COLORS['text_secondary'], pad=(0, 15))],
    *cal_header,
    *botones_calendario,
    cal_nav
], background_color=COLORS['surface'], pad=(15, 15), vertical_alignment='top', element_justification='center')

# --- Columna central (tabla) ---
data_col = sg.Column(tabla_panel, background_color=COLORS['bg'], pad=(20, 20),
                     expand_x=True, expand_y=True)

# --- Columna derecha (formulario) ---
form_col = sg.Column(form_panel, background_color=COLORS['surface'], pad=(20, 20),
                     vertical_alignment='top', element_justification='center')

# --- Layout final ---
layout = [[sidebar_col, sg.VSeparator(color=COLORS['border']),
           data_col, sg.VSeparator(color=COLORS['border']),
           form_col]]

window = sg.Window('Inventory Flow', layout, size=(1400, 750), margins=(0, 0),
                   background_color=COLORS['bg'], resizable=True, finalize=True)

# ------------------------------------------------------------
# FUNCIÓN PARA REFRESCAR LA INTERFAZ (SÓLO ACTUALIZA, NO RECREA)
# ------------------------------------------------------------


def refrescar_ui():
    """Actualiza todos los elementos dinámicos de la UI."""
    # --- Actualizar botones de categoría ---
    for cat in ['VINO', 'LATAS', 'DULCES']:
        is_active = (cat == categoria_activa.upper())
        color_key = cat.lower().capitalize()  # 'Vino', 'Latas', 'Dulces'
        bg = CATEGORIAS[color_key]['color'] if is_active else '#18181B'
        txt = '#FFFFFF' if is_active else COLORS['text_secondary']
        window[f'-CAT-{cat}'].update(button_color=(txt, bg))

    # --- Actualizar títulos y colores ---
    window['-TITLE-'].update(f'📋 INVENTARIO DE {categoria_activa.upper()}')
    window['-CAT_LABEL-'].update(f'Categoría activa: {categoria_activa}',
                                 text_color=CATEGORIAS[categoria_activa]['color'])
    window['-GUARDAR-'].update(button_color=('#FFFFFF',
                               CATEGORIAS[categoria_activa]['color']))

    # --- Refrescar calendario (cambio de mes o día) ---
    actualizar_calendario()

    # --- Refrescar datos de tabla ---
    refrescar_tabla()


def actualizar_calendario():
    """Cambia el texto y color de los botones del calendario según el mes actual."""
    year, month = fecha_actual.year, fecha_actual.month
    cal = calendar.monthcalendar(year, month)

    # Actualizar título del calendario
    window['-CAL_TITLE-'].update(
        f"{calendar.month_name[month].upper()} {year}")

    # Para cada botón, actualizar texto y color
    for i in range(MAX_SEMANAS):
        for j in range(7):
            dia = 0
            if i < len(cal) and j < len(cal[i]):
                dia = cal[i][j]

            btn_key = f'-CALBTN-{i}-{j}'
            if dia != 0:
                fecha_btn = datetime(year, month, dia)
                # Determinar color
                if fecha_btn.date() == fecha_seleccionada.date():
                    bg = CATEGORIAS[categoria_activa]['color']
                    txt = '#FFFFFF'
                elif fecha_btn.date() == datetime.now().date():
                    bg = '#27272A'
                    txt = '#FACC15'
                else:
                    bg = '#18181B'
                    txt = COLORS['text_secondary']

                window[btn_key].update(text=str(dia), disabled=False,
                                       button_color=(txt, bg))
            else:
                # Día vacío: botón deshabilitado y sin texto
                window[btn_key].update(text='', disabled=True,
                                       button_color=(COLORS['text'], '#18181B'))


def refrescar_tabla(filtrar=True):
    """Carga los datos en la tabla según la fecha seleccionada."""
    fecha_str = fecha_seleccionada.strftime("%Y-%m-%d") if filtrar else None
    datos = obtener_registros(categoria_activa, fecha_str)
    window['-TABLE-'].update(values=datos)
    window['-TABLE-'].update(selected_row_colors=('#FFFFFF',
                             CATEGORIAS[categoria_activa]['color']))
    subtitulo = f'Vista del: {fecha_seleccionada.strftime("%d/%m/%Y")}' if filtrar else 'Vista: TODOS LOS REGISTROS'
    window['-SUBTITLE-'].update(subtitulo)


# ------------------------------------------------------------
# INICIALIZACIÓN DE LA UI
# ------------------------------------------------------------
refrescar_ui()

# ------------------------------------------------------------
# BUCLE PRINCIPAL DE EVENTOS
# ------------------------------------------------------------
while True:
    event, values = window.read()
    if event == sg.WINDOW_CLOSED:
        break

    # === CAMBIO DE CATEGORÍA ===
    if event == '-CAT-VINO':
        categoria_activa = 'Vino'
        refrescar_ui()
    elif event == '-CAT-LATAS':
        categoria_activa = 'Latas'
        refrescar_ui()
    elif event == '-CAT-DULCES':
        categoria_activa = 'Dulces'
        refrescar_ui()

    # === NAVEGACIÓN DEL CALENDARIO ===
    elif event == '-MES_ANT-':
        # Mes anterior
        mes = fecha_actual.month - 1
        anio = fecha_actual.year
        if mes < 1:
            mes = 12
            anio -= 1
        fecha_actual = datetime(anio, mes, 1)
        # Ajustar selección si es necesario
        if fecha_seleccionada.month != mes or fecha_seleccionada.year != anio:
            fecha_seleccionada = fecha_actual
        refrescar_ui()

    elif event == '-MES_SIG-':
        # Mes siguiente
        mes = fecha_actual.month + 1
        anio = fecha_actual.year
        if mes > 12:
            mes = 1
            anio += 1
        fecha_actual = datetime(anio, mes, 1)
        if fecha_seleccionada.month != mes or fecha_seleccionada.year != anio:
            fecha_seleccionada = fecha_actual
        refrescar_ui()

    elif event == '-HOY-':
        fecha_actual = datetime.now()
        fecha_seleccionada = fecha_actual
        window['-FECHA-'].update(fecha_seleccionada.strftime("%Y-%m-%d"))
        refrescar_ui()

    # === CLICK EN UN DÍA DEL CALENDARIO ===
    elif event.startswith('-CALBTN-'):
        # Extraer coordenadas i, j
        partes = event.replace('-CALBTN-', '').split('-')
        i, j = int(partes[0]), int(partes[1])
        year, month = fecha_actual.year, fecha_actual.month
        cal = calendar.monthcalendar(year, month)
        if i < len(cal) and j < len(cal[i]):
            dia = cal[i][j]
            if dia != 0:
                fecha_seleccionada = datetime(year, month, dia)
                window['-FECHA-'].update(fecha_seleccionada.strftime("%Y-%m-%d"))
                refrescar_ui()

    # === GUARDAR PRODUCTO ===
    elif event == '-GUARDAR-':
        nombre = values['-NOMBRE-'].strip()
        cantidad = values['-CANTIDAD-'].strip()
        if not nombre or not cantidad:
            window['-MSG-'].update('❌ Campos vacíos',
                                   text_color=COLORS['danger'])
        elif not cantidad.isdigit():
            window['-MSG-'].update('❌ Cantidad inválida',
                                   text_color=COLORS['danger'])
        else:
            exito, msg = guardar_registro(
                categoria_activa, nombre, int(cantidad))
            if exito:
                window['-MSG-'].update(f'✅ {nombre} añadido',
                                       text_color=COLORS['success'])
                window['-NOMBRE-'].update('')
                window['-CANTIDAD-'].update('')
                refrescar_tabla()
                window['-NOMBRE-'].set_focus()
            else:
                window['-MSG-'].update(f'❌ {msg}', text_color=COLORS['danger'])

    # --- VER TODOS / ACTUALIZAR ---
    elif event == '-VER_TODOS-':
        refrescar_tabla(filtrar=False)
    elif event == '🔄 Actualizar':
        refrescar_tabla()
        window['-MSG-'].update('✅ Vista actualizada',
                               text_color=COLORS['success'])

    # --- ABRIR EXCEL ---
    elif event == '📂 Abrir Excel':
        if os.path.exists(ARCHIVO_EXCEL):
            os.startfile(ARCHIVO_EXCEL)

window.close()
