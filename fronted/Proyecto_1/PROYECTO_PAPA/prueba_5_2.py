import PySimpleGUI as sg
import pandas as pd
import os
from datetime import datetime
import calendar
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# --- Configuración de la ventana ---
sg.theme('LightBlue3')

# --- ÚNICO archivo Excel con PESTAÑAS separadas ---
ARCHIVO_EXCEL = 'registro_inventario.xlsx'
HOJAS_EXCEL = {
    'Vino': 'VINO',
    'Latas': 'LATAS',
    'Dulces': 'DULCES'
}

# --- Configuración de etiquetas dinámicas por categoría ---
ETIQUETAS_CATEGORIA = {
    'Vino': {
        'nombre': 'Nombre del vino:',
        'cantidad': 'Cantidad (botellas):',
        'titulo': '🍷 VINO',
        'color_titulo': '8B0000',
        'color_encabezado': 'CD5C5C'
    },
    'Latas': {
        'nombre': 'Nombre de la lata:',
        'cantidad': 'Cantidad (unidades):',
        'titulo': '🥫 LATAS',
        'color_titulo': '00008B',
        'color_encabezado': '4682B4'
    },
    'Dulces': {
        'nombre': 'Nombre del dulce:',
        'cantidad': 'Cantidad (piezas):',
        'titulo': '🍬 DULCES',
        'color_titulo': '006400',
        'color_encabezado': '3CB371'
    }
}

# --- Constante para altura de filas ---
ALTURA_FILA_DATOS = 25

# --- Función para inicializar el archivo Excel con las tres hojas ---


def inicializar_excel():
    """Crea el archivo Excel con tres hojas independientes"""
    if not os.path.exists(ARCHIVO_EXCEL):
        from openpyxl import Workbook
        wb = Workbook()

        # Eliminar la hoja por defecto "Sheet"
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])

        # Crear las tres hojas en orden
        for categoria, nombre_hoja in HOJAS_EXCEL.items():
            ws = wb.create_sheet(nombre_hoja)

            fila_actual = 1

            # Título de la tabla
            ws.merge_cells(f'A{fila_actual}:D{fila_actual}')
            celda_titulo = ws[f'A{fila_actual}']
            celda_titulo.value = ETIQUETAS_CATEGORIA[categoria]['titulo']
            celda_titulo.font = Font(size=14, bold=True, color='FFFFFF')
            celda_titulo.fill = PatternFill(
                start_color=ETIQUETAS_CATEGORIA[categoria]['color_titulo'],
                end_color=ETIQUETAS_CATEGORIA[categoria]['color_titulo'],
                fill_type='solid'
            )
            celda_titulo.alignment = Alignment(
                horizontal='center', vertical='center')
            ws.row_dimensions[fila_actual].height = 30

            fila_actual += 1

            # Encabezados
            encabezados = ['Nombre', 'Cantidad', 'Fecha', 'Hora']
            for col, encabezado in enumerate(encabezados, start=1):
                celda = ws.cell(row=fila_actual, column=col)
                celda.value = encabezado
                celda.font = Font(bold=True, color='FFFFFF', size=11)
                celda.fill = PatternFill(
                    start_color=ETIQUETAS_CATEGORIA[categoria]['color_encabezado'],
                    end_color=ETIQUETAS_CATEGORIA[categoria]['color_encabezado'],
                    fill_type='solid'
                )
                celda.alignment = Alignment(
                    horizontal='center', vertical='center')
                celda.border = Border(
                    left=Side(style='thin', color='000000'),
                    right=Side(style='thin', color='000000'),
                    top=Side(style='thin', color='000000'),
                    bottom=Side(style='thin', color='000000')
                )

            ws.row_dimensions[fila_actual].height = ALTURA_FILA_DATOS

            # Ajustar ancho de columnas
            ws.column_dimensions['A'].width = 35
            ws.column_dimensions['B'].width = 12
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 12

        wb.save(ARCHIVO_EXCEL)
        return True
    return False

# --- Función para guardar en la hoja correspondiente ---


def guardar_en_excel(categoria, nombre, cantidad, fecha=None, hora=None):
    """Añade un registro a la hoja específica"""
    if fecha is None:
        fecha = datetime.now().strftime("%Y-%m-%d")
    if hora is None:
        hora = datetime.now().strftime("%H:%M:%S")

    try:
        inicializar_excel()

        wb = load_workbook(ARCHIVO_EXCEL)
        ws = wb[HOJAS_EXCEL[categoria]]

        # Encontrar la primera fila vacía
        fila_datos = 3  # Empezamos después del título y encabezados

        while ws.cell(row=fila_datos, column=1).value is not None:
            fila_datos += 1
            if fila_datos > 1000:
                return False, "❌ Error: No se encontró espacio para guardar"

        # Establecer altura de la fila
        ws.row_dimensions[fila_datos].height = ALTURA_FILA_DATOS

        # Insertar datos
        ws.cell(row=fila_datos, column=1, value=nombre)
        ws.cell(row=fila_datos, column=2, value=cantidad)
        ws.cell(row=fila_datos, column=3, value=fecha)
        ws.cell(row=fila_datos, column=4, value=hora)

        # Aplicar estilo
        for col in range(1, 5):
            celda = ws.cell(row=fila_datos, column=col)
            celda.alignment = Alignment(
                horizontal='left', vertical='center', wrap_text=True)
            celda.border = Border(
                left=Side(style='thin', color='000000'),
                right=Side(style='thin', color='000000'),
                top=Side(style='thin', color='000000'),
                bottom=Side(style='thin', color='000000')
            )

        # Alternar colores
        if (fila_datos - 3) % 2 == 0:
            for col in range(1, 5):
                ws.cell(row=fila_datos, column=col).fill = PatternFill(
                    start_color='F5F5F5', end_color='F5F5F5', fill_type='solid'
                )

        wb.save(ARCHIVO_EXCEL)
        return True, f"✅ Registro guardado en {categoria}"

    except Exception as e:
        return False, f"❌ Error al guardar: {str(e)}"

# --- Función para filtrar registros por fecha ---


def filtrar_por_fecha(categoria, fecha_objetivo):
    """Obtiene los registros de una hoja específica filtrados por fecha"""
    if not os.path.exists(ARCHIVO_EXCEL):
        return []

    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)

        hoja_nombre = HOJAS_EXCEL[categoria]
        if hoja_nombre not in wb.sheetnames:
            wb.close()
            return []

        ws = wb[hoja_nombre]

        registros = []
        fecha_buscada = fecha_objetivo.strftime("%Y-%m-%d")

        # Leer desde la fila 3 (después de título y encabezados)
        fila_datos = 3

        while True:
            nombre = ws.cell(row=fila_datos, column=1).value
            if nombre is None:
                break

            cantidad = ws.cell(row=fila_datos, column=2).value
            fecha = ws.cell(row=fila_datos, column=3).value
            hora = ws.cell(row=fila_datos, column=4).value

            if fecha and fecha == fecha_buscada:
                registros.append(
                    [nombre, cantidad, fecha, hora if hora else ""])

            fila_datos += 1

        wb.close()

        registros.sort(key=lambda x: x[3] if x[3] else "", reverse=True)
        return registros

    except Exception as e:
        print(f"Error filtrando: {e}")
        return []

# --- Función MEJORADA para crear ventana de calendario con resaltado ---


def mostrar_calendario(fecha_actual, fecha_seleccionada_anterior=None):
    """
    Muestra un calendario interactivo con resaltado de colores:
    - Verde: Día actual
    - Amarillo/Dorado: Fecha seleccionada
    - Blanco: Días normales
    """
    año_actual = fecha_actual.year
    mes_actual = fecha_actual.month

    # Si no hay fecha seleccionada anterior, usar la fecha actual
    if fecha_seleccionada_anterior is None:
        fecha_seleccionada_anterior = fecha_actual

    cal = calendar.monthcalendar(año_actual, mes_actual)

    layout_calendario = [
        [sg.Text(f"{calendar.month_name[mes_actual]} {año_actual}",
                 font=('Arial', 16, 'bold'),
                 justification='center',
                 expand_x=True)],
        [sg.Text('')],

        [sg.Text(dia, size=(4, 1), justification='center', font=('Arial', 10, 'bold'))
         for dia in ['Lun', 'Mar', 'Mié', 'Jue', 'Vie', 'Sáb', 'Dom']],
    ]

    # Guardar referencias a los botones para poder actualizarlos
    botones_calendario = {}

    for semana in cal:
        fila_botones = []
        for dia in semana:
            if dia == 0:
                fila_botones.append(sg.Text('', size=(4, 2)))
            else:
                # Determinar el color del botón según el estado
                fecha_boton = datetime(año_actual, mes_actual, dia)

                # Prioridad de colores:
                # 1. Si es la fecha seleccionada -> Amarillo
                # 2. Si es el día actual -> Verde
                # 3. Día normal -> Blanco

                if (fecha_boton.date() == fecha_seleccionada_anterior.date()):
                    # Texto negro, fondo dorado
                    color_boton = ('black', '#FFD700')
                elif (dia == datetime.now().day and mes_actual == datetime.now().month
                      and año_actual == datetime.now().year):
                    color_boton = ('white', 'green')
                else:
                    color_boton = ('black', 'white')

                key_boton = f'-DIA_CAL_{dia}-'
                boton = sg.Button(str(dia),
                                  size=(4, 2),
                                  key=key_boton,
                                  button_color=color_boton)
                fila_botones.append(boton)
                botones_calendario[key_boton] = boton
        layout_calendario.append(fila_botones)

    layout_calendario.extend([
        [sg.Text('')],
        [sg.Button('◀ Mes Anterior', size=(12, 1), key='-MES_ANTERIOR-'),
         sg.Button('Hoy', size=(8, 1), key='-HOY-'),
         sg.Button('Mes Siguiente ▶', size=(12, 1), key='-MES_SIGUIENTE-')],
        [sg.Text('')],
        [sg.Text('', key='-INFO_SELECCION-', size=(30, 1), justification='center',
                 text_color='blue', font=('Arial', 9))],
        [sg.Text('')],
        [sg.Button('Cancelar', size=(10, 1)),
         sg.Button('Seleccionar', size=(10, 1), button_color=('white', 'blue'), key='-SELECCIONAR-')]
    ])

    window_cal = sg.Window('Seleccionar Fecha',
                           layout_calendario,
                           modal=True,
                           element_justification='center',
                           finalize=True)

    fecha_seleccionada = fecha_seleccionada_anterior
    mes_temporal = mes_actual
    año_temporal = año_actual

    # Actualizar información de selección
    if fecha_seleccionada:
        window_cal['-INFO_SELECCION-'].update(
            f"📌 Fecha seleccionada: {fecha_seleccionada.strftime('%d/%m/%Y')}"
        )

    while True:
        event_cal, values_cal = window_cal.read()

        if event_cal in (sg.WINDOW_CLOSED, 'Cancelar'):
            fecha_seleccionada = None
            break

        if event_cal == '-SELECCIONAR-':
            if fecha_seleccionada:
                break
            else:
                sg.popup('Por favor selecciona un día primero', title='Aviso')

        if event_cal == '-HOY-':
            fecha_seleccionada = datetime.now()
            break

        if event_cal == '-MES_ANTERIOR-':
            mes_temporal -= 1
            if mes_temporal < 1:
                mes_temporal = 12
                año_temporal -= 1
            window_cal.close()
            fecha_temporal = datetime(año_temporal, mes_temporal, 1)
            fecha_seleccionada = mostrar_calendario(
                fecha_temporal, fecha_seleccionada)
            break

        if event_cal == '-MES_SIGUIENTE-':
            mes_temporal += 1
            if mes_temporal > 12:
                mes_temporal = 1
                año_temporal += 1
            window_cal.close()
            fecha_temporal = datetime(año_temporal, mes_temporal, 1)
            fecha_seleccionada = mostrar_calendario(
                fecha_temporal, fecha_seleccionada)
            break

        # Manejar clic en un día del calendario
        if event_cal.startswith('-DIA_CAL_'):
            dia = int(event_cal.replace('-DIA_CAL_', '').replace('-', ''))
            nueva_fecha = datetime(año_temporal, mes_temporal, dia)

            # Actualizar la fecha seleccionada
            fecha_seleccionada = nueva_fecha

            # Resetear todos los botones a sus colores base
            for key, boton in botones_calendario.items():
                dia_boton = int(key.replace('-DIA_CAL_', '').replace('-', ''))
                fecha_boton = datetime(año_temporal, mes_temporal, dia_boton)

                if fecha_boton.date() == fecha_seleccionada.date():
                    # Seleccionado
                    boton.update(button_color=('black', '#FFD700'))
                elif (dia_boton == datetime.now().day and mes_temporal == datetime.now().month
                      and año_temporal == datetime.now().year):
                    boton.update(button_color=('white', 'green'))  # Día actual
                else:
                    boton.update(button_color=('black', 'white'))  # Normal

            # Actualizar información de selección
            window_cal['-INFO_SELECCION-'].update(
                f"📌 Fecha seleccionada: {fecha_seleccionada.strftime('%d/%m/%Y')}"
            )

    window_cal.close()
    return fecha_seleccionada


# --- Inicializar Excel ---
inicializar_excel()

# --- Fecha y hora actual ---
fecha_actual = datetime.now()
fecha_texto = fecha_actual.strftime("%d de %B de %Y")
hora_actual = fecha_actual.strftime("%H:%M:%S")

# --- Categoría actual ---
categoria_actual = 'Vino'

# --- Diseño de la interfaz ---
layout = [
    # === SELECTOR DE CATEGORÍA ===
    [sg.Text('📦 SELECCIONAR CATEGORÍA (Pestaña en Excel)', font=('Arial', 14, 'bold'),
             justification='center', expand_x=True)],

    [sg.Button('🍷 VINO', size=(15, 2), font=('Arial', 12, 'bold'),
               button_color=('white', 'darkred'), key='-CAT_VINO-'),
     sg.Button('🥫 LATAS', size=(15, 2), font=('Arial', 12, 'bold'),
               button_color=('white', 'darkblue'), key='-CAT_LATAS-'),
     sg.Button('🍬 DULCES', size=(15, 2), font=('Arial', 12, 'bold'),
               button_color=('white', 'darkgreen'), key='-CAT_DULCES-')],

    [sg.Text(f'📍 Pestaña activa: {HOJAS_EXCEL[categoria_actual]}',
             key='-PESTANA_ACTIVA-', font=('Arial', 11, 'italic'),
             text_color='purple', expand_x=True)],

    [sg.HorizontalSeparator()],

    # === SELECTOR DE FECHA ===
    [sg.Text('📅 FILTRAR POR FECHA', font=('Arial', 12, 'bold'),
             justification='center', expand_x=True)],

    [sg.Button('📆 Seleccionar Fecha', size=(20, 2), font=('Arial', 11, 'bold'),
               button_color=('white', 'purple'), key='-SELECCIONAR_FECHA-'),
     sg.Text(fecha_texto, key='-FECHA_MOSTRADA-', font=('Arial', 14, 'bold'),
             size=(30, 1), justification='center')],

    [sg.HorizontalSeparator()],
    [sg.Text('')],

    # === CONTENIDO PRINCIPAL ===
    [
        # COLUMNA IZQUIERDA: Tabla de registros
        sg.Column([
            [sg.Text('📋 REGISTROS DEL DÍA', font=('Arial', 12, 'bold'),
                     key='-TITULO_TABLA-')],
            [sg.Table(
                values=[],
                headings=['Nombre', 'Cantidad', 'Fecha', 'Hora'],
                max_col_width=30,
                auto_size_columns=False,
                col_widths=[30, 12, 15, 10],
                display_row_numbers=True,
                justification='left',
                num_rows=20,
                key='-TABLA-',
                row_height=25,
                enable_events=True,
                expand_x=True,
                expand_y=True
            )],
            [sg.Button('🔄 Actualizar', size=(15, 1)),
             sg.Button('📂 Abrir Excel', size=(15, 1),
                       button_color=('white', 'green')),
             sg.Button('📊 Ver Todos', size=(15, 1), button_color=('white', 'orange'))]
        ], expand_x=True, expand_y=True, scrollable=True),

        # COLUMNA DERECHA: Formulario
        sg.Column([
            [sg.Text(f'✨ NUEVO REGISTRO - {categoria_actual.upper()}',
                     font=('Arial', 12, 'bold'), key='-TITULO_FORMULARIO-')],
            [sg.Text('')],

            [sg.Text(ETIQUETAS_CATEGORIA[categoria_actual]['nombre'],
                     size=(20, 1), key='-LABEL_NOMBRE-', font=('Arial', 11))],
            [sg.InputText(key='-NOMBRE-', size=(30, 1), font=('Arial', 11))],

            [sg.Text('')],

            [sg.Text(ETIQUETAS_CATEGORIA[categoria_actual]['cantidad'],
                     size=(20, 1), key='-LABEL_CANTIDAD-', font=('Arial', 11))],
            [sg.InputText(key='-CANTIDAD-', size=(30, 1), font=('Arial', 11))],

            [sg.Text('')],

            [sg.Text('Fecha de registro:', size=(20, 1))],
            [sg.InputText(key='-FECHA_REGISTRO-', size=(30, 1),
                          default_text=fecha_actual.strftime("%Y-%m-%d"),
                          font=('Arial', 10))],

            [sg.Text('')],

            [sg.Text('Hora de registro:', size=(20, 1))],
            [sg.InputText(key='-HORA_REGISTRO-', size=(30, 1),
                          default_text=hora_actual,
                          font=('Arial', 10))],

            [sg.Text('')],
            [sg.Text('')],

            [sg.Button('💾 Guardar Registro', size=(20, 2),
                       button_color=('white', 'blue'),
                       font=('Arial', 11, 'bold'),
                       bind_return_key=True,
                       key='-GUARDAR-')],

            [sg.Text('')],

            [sg.Text('', key='-MENSAJE-', size=(40, 2), text_color='green',
                     font=('Arial', 10), justification='center')],

            [sg.Text('')],
            [sg.Button('🗑️ Limpiar campos', size=(20, 1),
                       button_color=('white', 'gray'))],
            [sg.Button('🕐 Usar hora actual', size=(20, 1),
                       button_color=('white', 'blue'), key='-HORA_ACTUAL-')]
        ], element_justification='center', vertical_alignment='top')
    ],

    # === BARRA DE ESTADO ===
    [sg.HorizontalSeparator()],
    [sg.Text('💡 Tip: El calendario resalta en AMARILLO la fecha seleccionada',
             text_color='gray', font=('Arial', 9), expand_x=True)]
]

# --- Crear ventana principal ---
window = sg.Window('Tienda Registro - Calendario con Resaltado',
                   layout,
                   size=(1100, 680),
                   resizable=True,
                   finalize=True)

# --- Variables globales ---
fecha_seleccionada_global = fecha_actual

# --- Función para actualizar tabla ---


def actualizar_tabla_por_fecha(categoria, fecha_filtro):
    registros = filtrar_por_fecha(categoria, fecha_filtro)
    window['-TABLA-'].update(values=registros)

    fecha_formateada = fecha_filtro.strftime("%d/%m/%Y")
    window['-TITULO_TABLA-'].update(
        f'📋 {categoria.upper()} - Registros del {fecha_formateada}')

    return len(registros) > 0

# --- Función para ver todos los registros ---


def ver_todos_registros(categoria):
    if not os.path.exists(ARCHIVO_EXCEL):
        return False

    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        ws = wb[HOJAS_EXCEL[categoria]]

        registros = []
        fila_datos = 3

        while True:
            nombre = ws.cell(row=fila_datos, column=1).value
            if nombre is None:
                break

            cantidad = ws.cell(row=fila_datos, column=2).value
            fecha = ws.cell(row=fila_datos, column=3).value
            hora = ws.cell(row=fila_datos, column=4).value
            registros.append([nombre, cantidad, fecha, hora if hora else ""])
            fila_datos += 1

        wb.close()

        if registros:
            registros.sort(key=lambda x: f"{x[2]} {x[3]}", reverse=True)
            window['-TABLA-'].update(values=registros)
            window['-TITULO_TABLA-'].update(
                f'📋 {categoria.upper()} - TODOS LOS REGISTROS')
            return True

    except Exception as e:
        sg.popup_error(f"Error al cargar: {e}")
    return False

# --- Función para cambiar de categoría ---


def cambiar_categoria(nueva_categoria):
    global categoria_actual
    categoria_actual = nueva_categoria

    window['-LABEL_NOMBRE-'].update(
        ETIQUETAS_CATEGORIA[categoria_actual]['nombre'])
    window['-LABEL_CANTIDAD-'].update(
        ETIQUETAS_CATEGORIA[categoria_actual]['cantidad'])
    window['-TITULO_FORMULARIO-'].update(
        f'✨ NUEVO REGISTRO - {categoria_actual.upper()}')
    window['-PESTANA_ACTIVA-'].update(
        f'📍 Pestaña activa: {HOJAS_EXCEL[categoria_actual]}')

    actualizar_tabla_por_fecha(categoria_actual, fecha_seleccionada_global)
    window['-MENSAJE-'].update(
        f'✅ Categoría cambiada a {categoria_actual}', text_color='green')


# --- Cargar registros iniciales ---
actualizar_tabla_por_fecha(categoria_actual, fecha_actual)

# --- Bucle principal ---
while True:
    event, values = window.read()

    if event == sg.WINDOW_CLOSED:
        break

    window['-FECHA_MOSTRADA-'].update(
        fecha_seleccionada_global.strftime("%d de %B de %Y"))

    # Cambiar categoría
    if event == '-CAT_VINO-':
        cambiar_categoria('Vino')
    elif event == '-CAT_LATAS-':
        cambiar_categoria('Latas')
    elif event == '-CAT_DULCES-':
        cambiar_categoria('Dulces')

    # Seleccionar fecha
    if event == '-SELECCIONAR_FECHA-':
        fecha_elegida = mostrar_calendario(
            fecha_actual, fecha_seleccionada_global)
        if fecha_elegida:
            fecha_seleccionada_global = fecha_elegida
            window['-FECHA_REGISTRO-'].update(
                fecha_elegida.strftime("%Y-%m-%d"))
            actualizar_tabla_por_fecha(
                categoria_actual, fecha_seleccionada_global)
            window['-MENSAJE-'].update(f"✅ Mostrando registros del {fecha_elegida.strftime('%d/%m/%Y')}",
                                       text_color='green')

    # Hora actual
    if event == '-HORA_ACTUAL-':
        hora_actual = datetime.now().strftime("%H:%M:%S")
        window['-HORA_REGISTRO-'].update(hora_actual)
        window['-MENSAJE-'].update("🕐 Hora actual actualizada",
                                   text_color='blue')

    # Guardar registro
    if event == '-GUARDAR-':
        nombre = values['-NOMBRE-'].strip()
        cantidad = values['-CANTIDAD-'].strip()
        fecha = values['-FECHA_REGISTRO-'].strip()
        hora = values['-HORA_REGISTRO-'].strip()

        if not nombre or not cantidad or not fecha or not hora:
            window['-MENSAJE-'].update(
                "❌ Todos los campos son obligatorios", text_color='red')
        elif not cantidad.isdigit():
            window['-MENSAJE-'].update("❌ La cantidad debe ser un número",
                                       text_color='red')
        else:
            exito, mensaje = guardar_en_excel(
                categoria_actual, nombre, int(cantidad), fecha, hora)

            if exito:
                window['-MENSAJE-'].update(mensaje, text_color='green')
                window['-NOMBRE-'].update('')
                window['-CANTIDAD-'].update('')
                actualizar_tabla_por_fecha(
                    categoria_actual, fecha_seleccionada_global)
                window['-NOMBRE-'].set_focus()
            else:
                window['-MENSAJE-'].update(mensaje, text_color='red')

    # Actualizar tabla
    if event == '🔄 Actualizar':
        actualizar_tabla_por_fecha(categoria_actual, fecha_seleccionada_global)
        window['-MENSAJE-'].update("✅ Tabla actualizada", text_color='green')

    # Ver todos
    if event == '📊 Ver Todos':
        if ver_todos_registros(categoria_actual):
            window['-MENSAJE-'].update(f"✅ Mostrando todos los registros de {categoria_actual}",
                                       text_color='green')
        else:
            window['-MENSAJE-'].update(f"ℹ️ No hay registros en {categoria_actual}",
                                       text_color='orange')

    # Abrir Excel
    if event == '📂 Abrir Excel':
        if os.path.exists(ARCHIVO_EXCEL):
            os.startfile(ARCHIVO_EXCEL)
        else:
            sg.popup("📁 Aún no hay archivo Excel. Guarda tu primer registro.",
                     title="Información")

    # Limpiar campos
    if event == '🗑️ Limpiar campos':
        window['-NOMBRE-'].update('')
        window['-CANTIDAD-'].update('')
        window['-MENSAJE-'].update("🧹 Campos limpiados", text_color='blue')
        window['-NOMBRE-'].set_focus()

window.close()
