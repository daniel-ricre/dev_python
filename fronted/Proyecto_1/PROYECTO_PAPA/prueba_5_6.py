import PySimpleGUI as sg
import pandas as pd
import os
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# --- Configuración básica ---
sg.theme('LightBlue2')
ARCHIVO_EXCEL = 'registro_inventario.xlsx'
HOJAS_EXCEL = {'Vino': 'VINO', 'Latas': 'LATAS', 'Dulces': 'DULCES'}

# --- Función simple para inicializar Excel ---


def inicializar_excel():
    if not os.path.exists(ARCHIVO_EXCEL):
        wb = Workbook()
        wb.remove(wb['Sheet'])
        for nombre_hoja in HOJAS_EXCEL.values():
            ws = wb.create_sheet(nombre_hoja)
            ws['A1'] = 'Nombre'
            ws['B1'] = 'Cantidad'
            ws['C1'] = 'Fecha'
            ws['D1'] = 'Hora'
        wb.save(ARCHIVO_EXCEL)

# --- Función simple para guardar ---


def guardar_en_excel(categoria, nombre, cantidad):
    try:
        inicializar_excel()
        fecha = datetime.now().strftime("%Y-%m-%d")
        hora = datetime.now().strftime("%H:%M:%S")

        wb = load_workbook(ARCHIVO_EXCEL)
        ws = wb[HOJAS_EXCEL[categoria]]
        fila = ws.max_row + 1
        ws.cell(row=fila, column=1, value=nombre)
        ws.cell(row=fila, column=2, value=cantidad)
        ws.cell(row=fila, column=3, value=fecha)
        ws.cell(row=fila, column=4, value=hora)
        wb.save(ARCHIVO_EXCEL)
        return True, "Guardado correctamente"
    except Exception as e:
        return False, str(e)

# --- Función simple para leer ---


def leer_registros(categoria):
    if not os.path.exists(ARCHIVO_EXCEL):
        return []
    try:
        df = pd.read_excel(ARCHIVO_EXCEL, sheet_name=HOJAS_EXCEL[categoria])
        return df.values.tolist()
    except:
        return []


# --- Interfaz mínima ---
categoria_actual = 'Vino'

layout = [
    [sg.Text('SISTEMA DE INVENTARIO', font=('Arial', 16, 'bold'),
             expand_x=True, justification='center')],
    [sg.Text('')],

    # Botones de categoría
    [sg.Button('🍷 VINO', size=(12, 2), key='-VINO-', button_color=('white', '#8B0000')),
     sg.Button('🥫 LATAS', size=(12, 2), key='-LATAS-',
               button_color=('white', '#00008B')),
     sg.Button('🍬 DULCES', size=(12, 2), key='-DULCES-', button_color=('white', '#006400'))],

    [sg.Text('')],
    [sg.Text(f'Categoría: {categoria_actual}',
             key='-CAT_TEXTO-', font=('Arial', 11, 'bold'))],
    [sg.Text('')],

    # Formulario
    [sg.Text('Nombre del producto:')],
    [sg.Input(key='-NOMBRE-', size=(40, 1))],
    [sg.Text('Cantidad:')],
    [sg.Input(key='-CANTIDAD-', size=(40, 1))],
    [sg.Button('GUARDAR', size=(15, 1), key='-GUARDAR-',
               button_color=('white', 'green'))],

    [sg.Text('')],
    [sg.Text('', key='-MENSAJE-', size=(50, 2), text_color='green')],
    [sg.Text('')],

    # Tabla
    [sg.Text('REGISTROS:', font=('Arial', 11, 'bold'))],
    [sg.Table(values=[], headings=['Nombre', 'Cantidad', 'Fecha', 'Hora'],
              col_widths=[25, 10, 15, 12], num_rows=12, key='-TABLA-',
              expand_x=True, expand_y=True)],

    [sg.Button('Actualizar', size=(12, 1)),
     sg.Button('Abrir Excel', size=(12, 1)),
     sg.Button('Salir', size=(12, 1))]
]

# Crear ventana
window = sg.Window('Inventario', layout, size=(
    700, 600), resizable=True, finalize=True)

# Función para actualizar tabla


def actualizar_tabla():
    datos = leer_registros(categoria_actual)
    window['-TABLA-'].update(values=datos)


# Cargar datos iniciales
actualizar_tabla()

# Bucle principal
while True:
    event, values = window.read()

    if event in (sg.WINDOW_CLOSED, 'Salir'):
        break

    # Cambiar categoría
    if event == '-VINO-':
        categoria_actual = 'Vino'
        window['-CAT_TEXTO-'].update('Categoría: Vino')
        actualizar_tabla()
    elif event == '-LATAS-':
        categoria_actual = 'Latas'
        window['-CAT_TEXTO-'].update('Categoría: Latas')
        actualizar_tabla()
    elif event == '-DULCES-':
        categoria_actual = 'Dulces'
        window['-CAT_TEXTO-'].update('Categoría: Dulces')
        actualizar_tabla()

    # Guardar
    if event == '-GUARDAR-':
        nombre = values['-NOMBRE-'].strip()
        cantidad = values['-CANTIDAD-'].strip()

        if not nombre or not cantidad:
            window['-MENSAJE-'].update('❌ Complete todos los campos',
                                       text_color='red')
        elif not cantidad.isdigit():
            window['-MENSAJE-'].update('❌ La cantidad debe ser un número',
                                       text_color='red')
        else:
            exito, msg = guardar_en_excel(
                categoria_actual, nombre, int(cantidad))
            if exito:
                window['-MENSAJE-'].update(f'✅ {msg}', text_color='green')
                window['-NOMBRE-'].update('')
                window['-CANTIDAD-'].update('')
                actualizar_tabla()
                window['-NOMBRE-'].set_focus()
            else:
                window['-MENSAJE-'].update(f'❌ Error: {msg}', text_color='red')

    # Actualizar tabla
    if event == 'Actualizar':
        actualizar_tabla()
        window['-MENSAJE-'].update('✅ Tabla actualizada', text_color='green')

    # Abrir Excel
    if event == 'Abrir Excel':
        if os.path.exists(ARCHIVO_EXCEL):
            os.startfile(ARCHIVO_EXCEL)
        else:
            window['-MENSAJE-'].update('⚠️ Aún no hay archivo Excel',
                                       text_color='orange')

window.close()
